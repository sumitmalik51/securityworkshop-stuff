#!/usr/bin/env python3
"""
Merge all batch CSV results with spns.xlsx credentials and create a consolidated
Excel workbook with separate sheets per status (COMPLETED, RUNNING, NOT_FOUND, ERROR).
"""

import os
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

BATCHES = ["Batch1", "Batch2", "Batch3", "Batch4", "Batch5"]
CSV_DIR = r"C:\certs"
SPNS_FILE = r"C:\certs\spns.xlsx"
OUTPUT_FILE = r"C:\certs\dspm_consolidated_status.xlsx"

# Column headers for output
HEADERS = [
    "Batch", "ODL User", "ODL Password", "TAP",
    "Tenant ID", "Subscription ID", "App ID", "App Secret",
    "Org Name", "Lab Name",
    "Status", "Job Title", "Items Found", "Job Date", "Checked At",
    "Screenshot File", "Remark"
]

# Remarks from screenshot analysis for NOT_FOUND accounts
NOT_FOUND_REMARKS = {
    # Batch1
    "odl_user_2226984": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2226992": "Switch to classic portal dialog blocking content",
    "odl_user_2227021": "Switch to classic portal dialog blocking content",
    "odl_user_2227233": "Switch to classic portal dialog blocking content",
    "odl_user_2227842": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227843": "Switch to classic portal dialog blocking content",
    # Batch2
    "odl_user_2227039": "Login failed - stuck on Microsoft Sign-in page",
    "odl_user_2227043": "Blank/empty Posture Agent tab",
    "odl_user_2227059": "Blank/empty Posture Agent tab",
    "odl_user_2227242": "Blank/empty Posture Agent tab",
    "odl_user_2227844": "Blank/empty Posture Agent tab",
    "odl_user_2227845": "Insufficient permissions - needs Purview Content Analyst role",
    # Batch3
    "odl_user_2227091": "Blank/empty Asset explorer page",
    "odl_user_2227093": "Blank/empty Asset explorer page",
    "odl_user_2227126": "Blank/empty Asset explorer page",
    "odl_user_2227129": "Page loading timeout - spinner/shimmer visible",
    "odl_user_2227131": "Blank/empty Asset explorer page",
    # Batch4
    "odl_user_2227168": "Blank/empty Asset explorer page",
    "odl_user_2227178": "Blank/empty Asset explorer page",
    "odl_user_2227846": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227848": "Blank/empty Asset explorer page",
    "odl_user_2227855": "Blank/empty Asset explorer page",
    "odl_user_2227856": "Insufficient permissions - needs Purview Content Analyst role",
    # Batch5
    "odl_user_2227195": "Blank/empty Posture Agent tab",
    "odl_user_2227196": "Blank/empty Posture Agent tab",
    "odl_user_2227321": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227322": "Blank/empty Posture Agent tab",
    "odl_user_2227323": "Blank/empty Posture Agent tab",
    "odl_user_2227324": "Blank/empty Posture Agent tab",
    "odl_user_2227325": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227326": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227327": "Blank/empty Posture Agent tab",
    "odl_user_2227328": "Blank/empty Posture Agent tab",
    "odl_user_2227330": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227331": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227332": "Blank/empty Posture Agent tab",
    "odl_user_2227335": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227337": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227339": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227342": "Blank/empty Posture Agent tab",
    "odl_user_2227343": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227345": "Blank/empty Posture Agent tab",
    "odl_user_2227346": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227414": "Blank/empty Posture Agent tab",
    "odl_user_2227418": "Blank/empty Posture Agent tab",
    "odl_user_2227421": "Blank/empty Posture Agent tab",
    "odl_user_2227423": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227495": "Blank/empty Posture Agent tab",
    "odl_user_2227496": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227498": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227500": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227501": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227502": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227503": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227506": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227507": "Blank/empty Posture Agent tab",
    "odl_user_2227508": "Blank/empty Posture Agent tab",
    "odl_user_2227510": "Blank/empty Posture Agent tab",
    "odl_user_2227511": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227512": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227513": "Blank/empty Posture Agent tab",
    "odl_user_2227812": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227813": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227821": "Blank/empty Posture Agent tab",
    "odl_user_2227825": "Insufficient permissions - needs Purview Content Analyst role",
    "odl_user_2227827": "Insufficient permissions - needs Purview Content Analyst role",
}


def load_spns_data():
    """Load all credentials from spns.xlsx into a dict keyed by odluser email."""
    wb = openpyxl.load_workbook(SPNS_FILE, read_only=True, data_only=True)
    spns = {}
    for batch in BATCHES:
        if batch not in wb.sheetnames:
            print(f"  [WARN] Sheet '{batch}' not found in spns.xlsx")
            continue
        ws = wb[batch]
        headers_row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        col_map = {h: i for i, h in enumerate(headers_row)}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[col_map.get("odluser", 6)]:
                continue
            odluser = str(row[col_map.get("odluser", 6)]).strip()
            spns[odluser] = {
                "batch": batch,
                "lab_name": str(row[col_map.get("Name", 0)] or ""),
                "subscription_id": str(row[col_map.get("SubscriptionId", 1)] or ""),
                "tenant_id": str(row[col_map.get("TenantId", 2)] or ""),
                "app_id": str(row[col_map.get("AppId", 3)] or ""),
                "app_secret": str(row[col_map.get("AppSecret", 4)] or ""),
                "orgname": str(row[col_map.get("orgname", 5)] or ""),
                "odluser": odluser,
                "odlpassword": str(row[col_map.get("odlpassword", 7)] or ""),
                "tap": str(row[col_map.get("TAP", 8)] or ""),
            }
    wb.close()
    print(f"Loaded {len(spns)} accounts from spns.xlsx")
    return spns


def load_csv_results():
    """Load all batch CSV results into a list of dicts."""
    results = []
    for batch in BATCHES:
        csv_file = os.path.join(CSV_DIR, f"dspm_status_{batch}.csv")
        if not os.path.exists(csv_file):
            print(f"  [WARN] {csv_file} not found, skipping")
            continue
        with open(csv_file, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            count = 0
            for row in reader:
                row["batch"] = batch
                results.append(row)
                count += 1
        print(f"Loaded {count} results from {batch}")
    print(f"Total results: {len(results)}")
    return results


def create_excel(results, spns):
    """Create the consolidated Excel with separate sheets per status."""
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Group results by status
    by_status = defaultdict(list)
    for r in results:
        by_status[r.get("status", "UNKNOWN")].append(r)

    # Define sheet order and colors
    sheet_config = {
        "COMPLETED": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),  # green
        "RUNNING": PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid"),    # blue
        "NOT_FOUND": PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),  # orange
        "ERROR": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),      # red
    }

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Also create a Summary sheet first
    summary_ws = wb.create_sheet("Summary")

    total_all = 0
    summary_data = []

    for status in ["COMPLETED", "RUNNING", "NOT_FOUND", "ERROR"]:
        rows = by_status.get(status, [])
        if not rows and status not in ["ERROR"]:
            # Create sheet even if empty
            pass

        ws = wb.create_sheet(status)
        fill = sheet_config.get(status, PatternFill())

        # Write headers
        for col_idx, header in enumerate(HEADERS, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = thin_border

        # Write data rows
        for row_idx, r in enumerate(rows, 2):
            username = r.get("username", "")
            spn = spns.get(username, {})

            # Derive remark for NOT_FOUND accounts from screenshot analysis
            remark = ""
            if status == "NOT_FOUND":
                # Extract odl_user key from username (email) e.g. odl_user_2227195@...
                user_key = username.split("@")[0] if "@" in username else username
                remark = NOT_FOUND_REMARKS.get(user_key, "")

            row_data = [
                r.get("batch", ""),
                username,
                spn.get("odlpassword", ""),
                spn.get("tap", ""),
                r.get("tenant_id", ""),
                spn.get("subscription_id", ""),
                spn.get("app_id", ""),
                spn.get("app_secret", ""),
                spn.get("orgname", ""),
                spn.get("lab_name", ""),
                r.get("status", ""),
                r.get("job_title", ""),
                r.get("items_found", ""),
                r.get("job_date", ""),
                r.get("checked_at", ""),
                os.path.basename(r.get("screenshot", "")) if r.get("screenshot") else "",
                remark,
            ]

            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.fill = fill
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center", wrap_text=False)

        # Auto-size columns
        for col_idx in range(1, len(HEADERS) + 1):
            max_len = len(str(HEADERS[col_idx - 1]))
            for row_idx in range(2, min(len(rows) + 2, 20)):  # Sample first 18 rows
                val = ws.cell(row=row_idx, column=col_idx).value
                if val:
                    max_len = max(max_len, min(len(str(val)), 50))
            ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 3

        # Freeze header row
        ws.freeze_panes = "A2"
        # Add auto-filter
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

        count = len(rows)
        total_all += count
        summary_data.append((status, count))
        print(f"  {status}: {count} accounts")

    # Build Summary sheet
    summary_ws.cell(row=1, column=1, value="DSPM Status Consolidated Report").font = Font(bold=True, size=16)
    summary_ws.cell(row=2, column=1, value=f"Generated from Batch1-Batch5 check_dspm_status.py results")
    summary_ws.cell(row=3, column=1, value=f"Total accounts: {total_all}")

    # Summary table
    summary_ws.cell(row=5, column=1, value="Status").font = Font(bold=True)
    summary_ws.cell(row=5, column=2, value="Count").font = Font(bold=True)
    summary_ws.cell(row=5, column=3, value="Percentage").font = Font(bold=True)
    for i, (status, count) in enumerate(summary_data, 6):
        summary_ws.cell(row=i, column=1, value=status)
        summary_ws.cell(row=i, column=2, value=count)
        pct = f"{count/total_all*100:.1f}%" if total_all > 0 else "0%"
        summary_ws.cell(row=i, column=3, value=pct)
        fill = sheet_config.get(status)
        if fill:
            for c in range(1, 4):
                summary_ws.cell(row=i, column=c).fill = fill

    summary_ws.cell(row=6 + len(summary_data), column=1, value="TOTAL").font = Font(bold=True)
    summary_ws.cell(row=6 + len(summary_data), column=2, value=total_all).font = Font(bold=True)

    # Batch breakdown
    row_start = 6 + len(summary_data) + 2
    summary_ws.cell(row=row_start, column=1, value="Batch Breakdown").font = Font(bold=True, size=13)
    row_start += 1
    batch_headers = ["Batch", "Total", "COMPLETED", "RUNNING", "NOT_FOUND", "ERROR"]
    for ci, h in enumerate(batch_headers, 1):
        summary_ws.cell(row=row_start, column=ci, value=h).font = Font(bold=True)

    for bi, batch in enumerate(BATCHES, row_start + 1):
        batch_rows = [r for r in results if r.get("batch") == batch]
        summary_ws.cell(row=bi, column=1, value=batch)
        summary_ws.cell(row=bi, column=2, value=len(batch_rows))
        for si, status in enumerate(["COMPLETED", "RUNNING", "NOT_FOUND", "ERROR"], 3):
            cnt = len([r for r in batch_rows if r.get("status") == status])
            summary_ws.cell(row=bi, column=si, value=cnt)

    summary_ws.column_dimensions["A"].width = 20
    summary_ws.column_dimensions["B"].width = 12
    summary_ws.column_dimensions["C"].width = 12

    wb.save(OUTPUT_FILE)
    print(f"\nExcel saved to: {OUTPUT_FILE}")


def main():
    print("Loading spns.xlsx credentials...")
    spns = load_spns_data()

    print("\nLoading CSV results...")
    results = load_csv_results()

    print("\nCreating consolidated Excel...")
    create_excel(results, spns)


if __name__ == "__main__":
    main()
