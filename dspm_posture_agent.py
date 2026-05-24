import time
import sys
import os
import openpyxl
from concurrent.futures import ThreadPoolExecutor, as_completed
from playwright.sync_api import sync_playwright

SPN_FILE = r"C:\certs\spns.xlsx"
SPN_SHEET = sys.argv[1] if len(sys.argv) > 1 else "Batch1"
START_INDEX = int(sys.argv[2]) if len(sys.argv) > 2 else 0
MAX_COUNT = int(sys.argv[3]) if len(sys.argv) > 3 else 0
PARALLEL = int(sys.argv[4]) if len(sys.argv) > 4 else 1

def load_accounts():
    accounts = []
    wb = openpyxl.load_workbook(SPN_FILE, read_only=True)
    ws = wb[SPN_SHEET]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    user_col = headers.index("odluser")
    pwd_col = headers.index("odlpassword")
    tid_col = headers.index("TenantId")
    for row in ws.iter_rows(min_row=2, values_only=True):
        username = row[user_col]
        password = row[pwd_col]
        tenant_id = row[tid_col]
        if username and password and tenant_id:
            accounts.append((str(username).strip(), str(password).strip(), str(tenant_id).strip()))
    wb.close()
    return accounts

def close_popups(page):
    """Close welcome/tour/copilot popups."""
    # Get started
    try:
        btn = page.locator('button:has-text("Get started")')
        if btn.count() > 0 and btn.first.is_visible():
            btn.first.click(timeout=3000)
            time.sleep(2)
    except Exception:
        pass

    # Close/dismiss dialogs
    for _ in range(10):
        closed = False
        for sel in [
            'button[aria-label="Close"]', 'button[aria-label="Dismiss"]',
            'button:has-text("Close")', 'button:has-text("Dismiss")',
            'button:has-text("Got it")', 'button:has-text("Skip")',
            'button:has-text("OK")', 'button:has-text("No thanks")',
            'button:has-text("Maybe later")', 'button:has-text("Not now")',
            'div[role="dialog"] button[aria-label="Close"]',
        ]:
            try:
                btn = page.locator(sel)
                if btn.count() > 0 and btn.first.is_visible():
                    btn.first.click(timeout=2000)
                    closed = True
                    time.sleep(1)
            except Exception:
                pass
        if not closed:
            break

    # Close Copilot panel
    try:
        cp = page.locator('button[aria-label="Close Security Copilot"], button[aria-label="Close Copilot"], button[aria-label="Close panel"]')
        if cp.count() > 0 and cp.first.is_visible():
            cp.first.click(timeout=3000, force=True)
            time.sleep(1)
    except Exception:
        pass

    # Remove backdrop overlays
    try:
        page.evaluate('document.querySelectorAll(".fui-DialogSurface__backdrop").forEach(e => e.remove())')
    except Exception:
        pass


def handle_setup_dialogs(page):
    """Handle 'Get started' and 'Start setup' dialogs that may appear for fresh accounts."""
    # Step 1: "Get started" button
    clicked_get_started = False
    try:
        get_started = page.locator('button:has-text("Get started")')
        if get_started.count() > 0 and get_started.first.is_visible():
            get_started.first.click(timeout=5000)
            print("  Clicked Get started")
            clicked_get_started = True
            time.sleep(10)
    except Exception:
        pass
    # Step 2: "Start setup" button - wait for it to appear (especially after Get started)
    try:
        start_setup = page.locator('button:has-text("Start setup")')
        if clicked_get_started:
            # After Get started, wait up to 15s for Start setup to appear
            start_setup.first.wait_for(state="visible", timeout=15000)
        if start_setup.count() > 0 and start_setup.first.is_visible():
            start_setup.first.click(timeout=5000)
            print("  Clicked Start setup")
            time.sleep(12)
    except Exception:
        pass
    # Step 3: "Congrats" dialog with Close button (appears ~10s after Start setup)
    clicked_close = False
    try:
        close_btn = page.locator('button:has-text("Close")')
        if close_btn.count() > 0 and close_btn.first.is_visible():
            close_btn.first.click(timeout=5000)
            print("  Closed 'Congrats' setup dialog")
            clicked_close = True
            time.sleep(3)
    except Exception:
        pass
    # Step 4: Refresh page after setup completes so UI reloads properly
    if clicked_close or clicked_get_started:
        try:
            page.reload(wait_until="domcontentloaded", timeout=30000)
            print("  Refreshed page after setup")
            time.sleep(10)
        except Exception:
            pass
        return True
    return False


def process_tenant(username, password, tenant_id, index, total):
    print(f"\n=== [{index}/{total}] {username} ===")
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=False)
            context = browser.new_context()
            page = context.new_page()

            # Login
            page.goto("https://purview.microsoft.com")
            time.sleep(3)

            page.fill('input[type="email"]', username)
            page.click('input[type="submit"]')
            time.sleep(4)

            # Password option if needed
            try:
                pw_link = page.locator('a:has-text("Use your password instead"), a:has-text("password"), #idA_PWD_SwitchToPassword')
                pw_link.first.click(timeout=8000)
                time.sleep(3)
            except Exception:
                pass

            try:
                page.wait_for_selector('input[type="password"]', timeout=10000)
                page.fill('input[type="password"]', password)
                page.click('input[type="submit"]')
            except Exception as e:
                print(f"  [WARN] Password issue: {e}")
            time.sleep(4)

            # Stay signed in - Yes
            try:
                yes_btn = page.locator('input[type="submit"][value="Yes"], button:has-text("Yes")')
                yes_btn.first.click(timeout=15000)
            except Exception:
                pass
            time.sleep(5)

            try:
                page.wait_for_load_state("load", timeout=30000)
            except Exception:
                pass
            time.sleep(5)

            close_popups(page)
            time.sleep(2)

            # Navigate to DSPM Asset Explorer
            dspm_url = f"https://purview.microsoft.com/dspm/assetexplorer?tid={tenant_id}"
            print(f"  Navigating to: {dspm_url}")
            page.goto(dspm_url, wait_until="domcontentloaded", timeout=60000)
            time.sleep(10)
            close_popups(page)
            time.sleep(2)

            # Click "Posture agent (preview)" tab/option (use force=True to bypass overlays)
            print("  Clicking 'Posture agent (preview)'...")
            try:
                posture = page.locator('button:has-text("Posture Agent"), button:has-text("Posture agent")')
                if posture.count() == 0:
                    posture = page.get_by_text("Posture agent", exact=False)
                posture.first.click(timeout=10000, force=True)
                print("  Clicked Posture agent")
                time.sleep(8)
            except Exception as e:
                print(f"  [WARN] Could not click Posture agent: {e}")

            # Handle fresh account setup flow (can appear after clicking Posture Agent)
            did_setup = handle_setup_dialogs(page)

            close_popups(page)
            time.sleep(2)

            # If setup happened, page was refreshed - need to click Posture Agent again
            if did_setup:
                print("  Re-clicking 'Posture agent' after setup...")
                try:
                    posture = page.locator('button:has-text("Posture Agent"), button:has-text("Posture agent")')
                    if posture.count() == 0:
                        posture = page.get_by_text("Posture agent", exact=False)
                    posture.first.click(timeout=10000, force=True)
                    print("  Re-clicked Posture agent")
                    time.sleep(8)
                except Exception as e:
                    print(f"  [WARN] Could not re-click Posture agent: {e}")
                close_popups(page)
                time.sleep(2)

            # Click "Add data sources" button
            print("  Clicking 'Add data sources'...")
            try:
                # Remove any backdrop overlays first
                page.evaluate('document.querySelectorAll(".fui-DialogSurface__backdrop").forEach(e => e.remove())')
                time.sleep(1)
                add_data = page.locator('button:has-text("Add data sources"), a:has-text("Add data sources"), [role="button"]:has-text("Add data sources")')
                if add_data.count() == 0:
                    add_data = page.locator('button:has-text("Add data"), a:has-text("Add data"), [role="button"]:has-text("Add data")')
                if add_data.count() == 0:
                    add_data = page.get_by_text("Add data sources", exact=False)
                add_data.first.click(timeout=10000, force=True)
                print("  Clicked Add data sources")
                time.sleep(5)
            except Exception as e:
                print(f"  [WARN] Could not click Add data sources: {e}")
                try:
                    page.evaluate('''() => {
                        const walker = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT);
                        while (walker.nextNode()) {
                            if (walker.currentNode.textContent.trim().toLowerCase().includes('add data source')) {
                                walker.currentNode.parentElement.click();
                                return true;
                            }
                        }
                        return false;
                    }''')
                    print("  Clicked Add data sources via JS")
                    time.sleep(5)
                except Exception:
                    pass

            # Handle setup dialogs that can appear after clicking Add data sources
            did_setup2 = handle_setup_dialogs(page)

            # If setup happened after Add data sources, re-click Posture Agent + Add data sources
            if did_setup2:
                close_popups(page)
                time.sleep(2)
                print("  Re-clicking 'Posture agent' after late setup...")
                try:
                    posture = page.locator('button:has-text("Posture Agent"), button:has-text("Posture agent")')
                    if posture.count() == 0:
                        posture = page.get_by_text("Posture agent", exact=False)
                    posture.first.click(timeout=10000, force=True)
                    print("  Re-clicked Posture agent")
                    time.sleep(8)
                except Exception:
                    pass
                close_popups(page)
                time.sleep(2)
                print("  Re-clicking 'Add data sources'...")
                try:
                    add_data = page.locator('button:has-text("Add data sources"), a:has-text("Add data sources"), [role="button"]:has-text("Add data sources")')
                    if add_data.count() == 0:
                        add_data = page.get_by_text("Add data sources", exact=False)
                    add_data.first.click(timeout=10000)
                    print("  Re-clicked Add data sources")
                    time.sleep(5)
                except Exception:
                    pass

            close_popups(page)
            time.sleep(2)

            # Search for odluser username (just the user part)
            user_search = username.split("@")[0]
            print(f"  Searching for '{user_search}'...")
            try:
                # Use the data sources search box (NOT the global search at the top)
                search_input = page.locator('input[placeholder="Search for people, groups, or sites"]')
                search_input.fill(user_search)
                print(f"  Filled search: {user_search}")
                time.sleep(1)
                search_input.press("Enter")
                time.sleep(8)
            except Exception as e:
                print(f"  [WARN] Search issue: {e}")

            # Select the user from search results
            print("  Selecting user from results...")
            try:
                # Wait for results to appear, then click the matching button/row
                # The result is a button containing the username and email
                user_btn = page.locator(f'button:has-text("{user_search}")').last
                user_btn.wait_for(state="visible", timeout=15000)
                user_btn.click(timeout=5000)
                print(f"  Selected {user_search}")
                time.sleep(5)
            except Exception as e:
                print(f"  [WARN] Button click failed ({e}), trying alternatives...")
                try:
                    # Try clicking any element with the email text
                    email_el = page.locator(f'text="{username}"').last
                    if email_el.count() > 0:
                        email_el.click(timeout=5000)
                        print(f"  Selected via email text")
                        time.sleep(5)
                    else:
                        # JS fallback: find and click element with user text
                        page.evaluate(f'''() => {{
                            const els = document.querySelectorAll("button, div[role='option'], div[role='listitem'], li");
                            for (const el of els) {{
                                if (el.textContent.includes("{user_search}")) {{
                                    el.click();
                                    return true;
                                }}
                            }}
                            return false;
                        }}''')
                        print(f"  Selected via JS fallback")
                        time.sleep(5)
                except Exception as e2:
                    print(f"  [WARN] Could not select user: {e2}")

            # Handle setup dialog that can appear late (after search/select)
            handle_setup_dialogs(page)
            close_popups(page)

            # After selecting user, two header checkboxes appear: Mailboxes (cb[0]) and Sites (cb[1])
            # Both are checked by default. Need to uncheck Mailboxes, keep Sites.
            print("  Unchecking Mailboxes, keeping Sites...")
            try:
                # Wait for checkboxes to appear
                page.locator('input[type="checkbox"], [role="checkbox"]').first.wait_for(timeout=15000)
                time.sleep(2)
                
                checkboxes = page.locator('input[type="checkbox"], [role="checkbox"]')
                cb_count = checkboxes.count()
                print(f"  Found {cb_count} checkboxes")
                
                if cb_count >= 2:
                    # First checkbox is Mailboxes - uncheck via label or force click
                    mail_cb = checkboxes.nth(0)
                    if mail_cb.is_visible():
                        cb_id = mail_cb.get_attribute('id')
                        if cb_id:
                            label = page.locator(f'label[for="{cb_id}"]')
                            if label.count() > 0:
                                label.first.click(force=True)
                                print("  Unchecked Mailboxes via label")
                            else:
                                mail_cb.click(force=True)
                                print("  Unchecked Mailboxes (force)")
                        else:
                            mail_cb.click(force=True)
                            print("  Unchecked Mailboxes (force)")
                    # Second checkbox is Sites - leave it checked
                    print("  Sites left checked")
                elif cb_count == 1:
                    print("  [WARN] Only 1 checkbox found")
                else:
                    print("  [WARN] No checkboxes found after selecting user")
                time.sleep(2)
            except Exception as e:
                print(f"  [WARN] Checkbox issue: {e}")

            # Click Save button
            print("  Clicking Save...")
            try:
                save_btn = page.locator('button:has-text("Save")')
                save_btn.first.click(timeout=5000, force=True)
                print("  Clicked Save")
                time.sleep(8)
            except Exception as e:
                print(f"  [WARN] Could not click Save: {e}")

            # Wait 10 seconds after Save for page to settle
            print("  Waiting 10s after Save...")
            time.sleep(10)

            # Enter prompt and click Send
            print("  Entering prompt...")
            try:
                # Click "Start new prompt" if present (when previous job exists)
                try:
                    start_new = page.locator('button:has-text("Start new prompt")')
                    if start_new.count() > 0 and start_new.first.is_visible():
                        start_new.first.click(timeout=5000)
                        print("  Clicked Start new prompt")
                        time.sleep(5)
                except Exception:
                    pass

                # Click any suggested prompt button to activate the prompt area
                suggested_selectors = [
                    'button:has-text("Collect files")',
                    'button:has-text("Show chats")',
                    'button:has-text("Retrieve financial")',
                    'button:has-text("Suggested prompt")',
                ]
                clicked_suggested = False
                for sel in suggested_selectors:
                    try:
                        btn = page.locator(sel)
                        if btn.count() > 0 and btn.first.is_visible():
                            btn.first.click(timeout=5000)
                            print(f"  Clicked suggested prompt to activate area")
                            clicked_suggested = True
                            time.sleep(3)
                            break
                    except Exception:
                        continue

                if clicked_suggested:
                    # Now the prompt box should be active with the suggested text
                    # Select all and replace with correct prompt
                    page.keyboard.press("Control+a")
                    time.sleep(0.5)
                    page.keyboard.press("Delete")
                    time.sleep(0.5)
                    page.keyboard.type("Get files with sensitive types password and credentials", delay=50)
                    print("  Typed correct prompt")
                    time.sleep(3)
                else:
                    # Fallback: try clicking prompt area directly
                    prompt_box = page.get_by_placeholder("Describe what you're looking for")
                    prompt_box.click(timeout=10000)
                    time.sleep(1)
                    page.keyboard.type("Get files with sensitive types password and credentials", delay=50)
                    print("  Typed prompt directly")
                    time.sleep(3)

                # Click Send button (the arrow icon next to prompt)
                send_btn = page.locator('button[aria-label="Send"], button:has-text("Send")')
                if send_btn.count() > 0:
                    time.sleep(2)
                    send_btn.first.click(force=True, timeout=5000)
                    print("  Clicked Send")
                else:
                    page.keyboard.press("Enter")
                    print("  Pressed Enter to send")

                # Wait for confirmation page to load (can take up to 1 minute)
                print("  Waiting for Confirm page...")
                confirm_btn = page.locator('button:has-text("Confirm")')
                confirm_btn.wait_for(state="visible", timeout=90000)
                time.sleep(2)
                confirm_btn.click(timeout=5000, force=True)
                print("  Clicked Confirm")

                # Wait for job to start running, then refresh
                print("  Waiting 30s for job to start...")
                time.sleep(30)
                refresh_btn = page.locator('button[aria-label="Refresh"], button:has-text("Refresh"), button[title="Refresh"]')
                if refresh_btn.count() > 0:
                    refresh_btn.first.click(timeout=5000)
                    print("  Clicked Refresh")
                else:
                    page.reload()
                    print("  Reloaded page")
                print("  Waiting 30s before closing...")
                time.sleep(30)
            except Exception as e:
                print(f"  [WARN] Prompt/Confirm issue: {e}")

            print(f"  SUCCESS")
            time.sleep(2)
            browser.close()
            return True

    except Exception as e:
        print(f"  FAILED: {e}")
        return False


def main():
    accounts = load_accounts()
    print(f"Loaded {len(accounts)} accounts from sheet '{SPN_SHEET}'")

    # Apply start index and max count
    if START_INDEX > 0:
        accounts = accounts[START_INDEX:]
    if MAX_COUNT > 0:
        accounts = accounts[:MAX_COUNT]

    total = len(accounts)
    print(f"Processing {total} accounts (start={START_INDEX}, max={MAX_COUNT}, parallel={PARALLEL})")

    success = 0
    failed = 0

    if PARALLEL <= 1:
        # Sequential execution
        for i, (user, pwd, tid) in enumerate(accounts, 1):
            result = process_tenant(user, pwd, tid, i, total)
            if result:
                success += 1
            else:
                failed += 1
    else:
        # Parallel execution with ThreadPoolExecutor
        def run_task(args):
            idx, user, pwd, tid = args
            return (user, process_tenant(user, pwd, tid, idx, total))

        tasks = [(i, user, pwd, tid) for i, (user, pwd, tid) in enumerate(accounts, 1)]

        with ThreadPoolExecutor(max_workers=PARALLEL) as executor:
            futures = {executor.submit(run_task, task): task for task in tasks}
            for future in as_completed(futures):
                try:
                    user, result = future.result()
                    if result:
                        success += 1
                    else:
                        failed += 1
                except Exception as e:
                    failed += 1
                    print(f"  Thread error: {e}")
                print(f"  Progress: {success + failed}/{total} (Success={success}, Failed={failed})")

    print(f"\n=== COMPLETE: Success={success}, Failed={failed}, Total={total} ===")


if __name__ == "__main__":
    main()
