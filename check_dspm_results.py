#!/usr/bin/env python3
"""
Check DSPM Posture Agent job results by intercepting the DiscoveryAgent/discoveries API.
Much faster & more reliable than DOM scraping — captures ProcessingStatus, MatchesFound directly.

Usage:
  python check_dspm_results.py Batch1 0 0 5    # Batch1, start=0, count=0(all), parallel=5
  python check_dspm_results.py Batch2 0 10 3   # Batch2, first 10, 3 parallel
"""

import time
import sys
import os
import csv
import json
import openpyxl
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from playwright.sync_api import sync_playwright

SPN_FILE = r"C:\certs\spns.xlsx"
SPN_SHEET = sys.argv[1] if len(sys.argv) > 1 else "Batch1"
START_INDEX = int(sys.argv[2]) if len(sys.argv) > 2 else 0
MAX_COUNT = int(sys.argv[3]) if len(sys.argv) > 3 else 0
PARALLEL = int(sys.argv[4]) if len(sys.argv) > 4 else 5

RESULTS_FILE = rf"C:\certs\dspm_results_{SPN_SHEET}.csv"


def load_accounts():
    accounts = []
    wb = openpyxl.load_workbook(SPN_FILE, read_only=True)
    ws = wb[SPN_SHEET]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    user_col = headers.index("odluser")
    pwd_col = headers.index("odlpassword")
    tid_col = headers.index("TenantId")
    for row in ws.iter_rows(min_row=2, values_only=True):
        username = row[user_col]
        password = row[pwd_col]
        tenant_id = row[tid_col]
        if username and password and tenant_id:
            accounts.append((str(username).strip(), str(password).strip(), str(tenant_id).strip()))
    wb.close()
    return accounts


def close_popups(page):
    for _ in range(5):
        closed = False
        for sel in [
            'button[aria-label="Close"]', 'button[aria-label="Dismiss"]',
            'button:has-text("Close")', 'button:has-text("Dismiss")',
            'button:has-text("Got it")', 'button:has-text("Skip")',
            'button:has-text("OK")', 'button:has-text("No thanks")',
            'button:has-text("Maybe later")', 'button:has-text("Not now")',
            'div[role="dialog"] button[aria-label="Close"]',
        ]:
            try:
                btn = page.locator(sel)
                if btn.count() > 0 and btn.first.is_visible():
                    btn.first.click(timeout=2000)
                    closed = True
                    time.sleep(1)
            except Exception:
                pass
        if not closed:
            break
    try:
        page.evaluate('document.querySelectorAll(".fui-DialogSurface__backdrop").forEach(e => e.remove())')
    except Exception:
        pass


def handle_setup_dialogs(page):
    clicked_get_started = False
    try:
        gs = page.locator('button:has-text("Get started")')
        if gs.count() > 0 and gs.first.is_visible():
            gs.first.click(timeout=5000)
            clicked_get_started = True
            time.sleep(10)
    except Exception:
        pass
    try:
        ss = page.locator('button:has-text("Start setup")')
        if clicked_get_started:
            ss.first.wait_for(state="visible", timeout=15000)
        if ss.count() > 0 and ss.first.is_visible():
            ss.first.click(timeout=5000)
            time.sleep(12)
    except Exception:
        pass
    try:
        cb = page.locator('button:has-text("Close")')
        if cb.count() > 0 and cb.first.is_visible():
            cb.first.click(timeout=5000)
            time.sleep(3)
    except Exception:
        pass
    if clicked_get_started:
        try:
            page.reload(wait_until="domcontentloaded", timeout=120000)
            time.sleep(15)
        except Exception:
            pass


def check_tenant(username, password, tenant_id, index, total):
    print(f"\n=== [{index}/{total}] {username} ===")
    result = {
        "username": username,
        "tenant_id": tenant_id,
        "status": "ERROR",
        "matches_found": "",
        "processing_status": "",
        "discovery_name": "",
        "summary": "",
        "start_time": "",
        "end_time": "",
        "error_code": "",
        "checked_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }

    # Capture DiscoveryAgent/discoveries response via network interception
    discoveries_data = {}

    def on_response(response):
        nonlocal discoveries_data
        try:
            if "DiscoveryAgent/discoveries" in response.url and response.status == 200:
                body = response.json()
                discoveries_data["raw"] = body
        except Exception:
            pass

    for attempt in range(3):
        if attempt > 0:
            print(f"  Retry {attempt+1}/3")
            time.sleep(5)
        discoveries_data.clear()
        browser = None
        try:
            with sync_playwright() as p:
                browser = p.chromium.launch(headless=False)
                context = browser.new_context()
                page = context.new_page()

                # Listen for discoveries API response
                page.on("response", on_response)

                # Login
                for _nav in range(3):
                    try:
                        page.goto("https://purview.microsoft.com", timeout=120000)
                        break
                    except Exception:
                        if _nav == 2:
                            raise
                        time.sleep(10)
                time.sleep(3)
                try:
                    page.wait_for_load_state("load", timeout=30000)
                except Exception:
                    pass

                page.fill('input[type="email"]', username)
                page.click('input[type="submit"]')
                time.sleep(4)

                try:
                    pw_link = page.locator('a:has-text("Use your password instead"), a:has-text("password"), #idA_PWD_SwitchToPassword')
                    pw_link.first.click(timeout=8000)
                    time.sleep(3)
                except Exception:
                    pass

                page.fill('input[type="password"]', password)
                page.click('input[type="submit"]')
                time.sleep(5)

                try:
                    yes_btn = page.locator('input[type="submit"][value="Yes"], button:has-text("Yes")')
                    yes_btn.first.click(timeout=8000)
                    time.sleep(3)
                except Exception:
                    pass

                try:
                    page.wait_for_load_state("load", timeout=30000)
                except Exception:
                    pass
                time.sleep(3)
                close_popups(page)

                # Navigate to DSPM
                url = f"https://purview.microsoft.com/dspm/assetexplorer?tid={tenant_id}"
                print(f"  Navigating to DSPM...")
                for _nav in range(3):
                    try:
                        page.goto(url, wait_until="domcontentloaded", timeout=120000)
                        break
                    except Exception:
                        if _nav == 2:
                            raise
                        time.sleep(10)
                try:
                    page.wait_for_load_state("load", timeout=30000)
                except Exception:
                    pass
                time.sleep(15)

                # Handle ALL popups/setup dialogs first before looking for the tab
                for _popup_round in range(3):
                    close_popups(page)
                    handle_setup_dialogs(page)
                    close_popups(page)
                    time.sleep(3)
                    # Check if any dialog/overlay is still visible
                    try:
                        dialogs = page.locator('[role="dialog"], .ms-Dialog, .fui-DialogSurface')
                        if dialogs.count() == 0:
                            break
                        print(f"  Dialog still visible, clearing again...")
                    except Exception:
                        break

                # Click Posture Agent tab — use role="tab" to avoid matching popups
                pa_clicked = False
                for _w in range(36):
                    # Try specific tab selector first (role="tab")
                    for sel in [
                        'button[role="tab"]:has-text("Posture")',
                        '[role="tab"]:has-text("Posture")',
                        '[role="tablist"] button:has-text("Posture")',
                    ]:
                        try:
                            pa = page.locator(sel)
                            if pa.count() > 0 and pa.first.is_visible():
                                pa.first.click(timeout=5000, force=True)
                                print(f"  Clicked Posture Agent tab (via {sel})")
                                pa_clicked = True
                                break
                        except Exception:
                            pass
                    if pa_clicked:
                        break

                    # Fallback: try generic button but only if no dialogs are open
                    try:
                        dialogs = page.locator('[role="dialog"], .ms-Dialog, .fui-DialogSurface')
                        if dialogs.count() == 0:
                            pa = page.locator('button:has-text("Posture Agent"), button:has-text("Posture agent")')
                            if pa.count() > 0 and pa.first.is_visible():
                                pa.first.click(timeout=5000, force=True)
                                print("  Clicked Posture Agent tab (button fallback)")
                                pa_clicked = True
                                break
                    except Exception:
                        pass

                    # Still loading...
                    try:
                        shimmers = page.locator('.ms-Shimmer-container, .ms-Spinner, [class*="spinner"], [class*="loading"], [class*="Shimmer"]')
                        if shimmers.count() > 0 and _w % 6 == 0:
                            print(f"  Page still loading... ({_w*5}s)")
                    except Exception:
                        pass
                    close_popups(page)
                    handle_setup_dialogs(page)
                    time.sleep(5)

                if not pa_clicked:
                    print("  [WARN] Could not click Posture Agent after 180s")

                # Wait for Posture Agent content + discoveries API response (up to 90s)
                time.sleep(8)
                print("  Waiting for discoveries API...")
                for _w in range(18):
                    if "raw" in discoveries_data:
                        break
                    close_popups(page)
                    time.sleep(5)

                # Parse the API response
                if "raw" in discoveries_data:
                    data = discoveries_data["raw"]
                    values = data.get("Value") or data.get("value") or []
                    if values:
                        d = values[0]  # latest discovery
                        ps = d.get("ProcessingStatus", "")
                        mf = d.get("MatchesFound", 0)
                        result["processing_status"] = ps
                        result["matches_found"] = str(mf) if mf is not None else "0"
                        result["discovery_name"] = d.get("DiscoveryName", "")
                        result["summary"] = (d.get("Summary") or "")[:200]
                        result["start_time"] = d.get("StartExecutionTime", "")
                        result["end_time"] = d.get("EndExecutionTime", "")
                        result["error_code"] = str(d.get("ErrorCode", ""))

                        if ps == "Successful":
                            result["status"] = "COMPLETED"
                            print(f"  COMPLETED — {mf} files found")
                        elif ps in ("Running", "InProgress"):
                            result["status"] = "RUNNING"
                            print(f"  RUNNING")
                        elif ps in ("Failed", "Error"):
                            result["status"] = "FAILED"
                            print(f"  FAILED — {d.get('ErrorCode', '')}")
                        else:
                            result["status"] = ps or "UNKNOWN"
                            print(f"  Status: {ps}")
                    else:
                        result["status"] = "NO_DISCOVERY"
                        print("  NO_DISCOVERY — API returned empty list")
                else:
                    # Fallback: check DOM for "View insights" / running indicators
                    print("  [WARN] No API response captured, checking DOM...")
                    try:
                        vi = page.locator('button:has-text("View insights")')
                        if vi.count() > 0 and vi.first.is_visible():
                            result["status"] = "COMPLETED"
                            items = page.locator('text=/found \\d+ items/')
                            if items.count() > 0:
                                result["summary"] = items.first.inner_text()
                            print(f"  COMPLETED (DOM) — {result['summary']}")
                        elif page.locator('text=/[Rr]unning|[Ss]top generating/').count() > 0:
                            result["status"] = "RUNNING"
                            print("  RUNNING (DOM)")
                        else:
                            result["status"] = "NOT_FOUND"
                            print("  NOT_FOUND (DOM)")
                    except Exception:
                        result["status"] = "CHECK_ERROR"

        except Exception as e:
            print(f"  ERROR (attempt {attempt+1}): {e}")
            result["status"] = "ERROR"
        finally:
            try:
                if browser:
                    browser.close()
            except Exception:
                pass

        if result["status"] not in ("ERROR", "CHECK_ERROR"):
            break

    return result


def load_existing_results():
    existing = {}
    if os.path.exists(RESULTS_FILE):
        try:
            with open(RESULTS_FILE, "r", newline="", encoding="utf-8") as f:
                for row in csv.DictReader(f):
                    existing[row["username"]] = row
            completed = sum(1 for r in existing.values() if r["status"] == "COMPLETED")
            print(f"Loaded {len(existing)} existing results ({completed} COMPLETED, will skip)")
        except Exception:
            pass
    return existing


def save_results(all_results):
    fieldnames = ["username", "tenant_id", "status", "processing_status", "matches_found",
                  "discovery_name", "summary", "start_time", "end_time", "error_code", "checked_at"]
    with open(RESULTS_FILE, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(all_results.values())


def main():
    accounts = load_accounts()
    print(f"Loaded {len(accounts)} accounts from '{SPN_SHEET}'")

    all_results = load_existing_results()

    if START_INDEX > 0:
        accounts = accounts[START_INDEX:]
    if MAX_COUNT > 0:
        accounts = accounts[:MAX_COUNT]

    to_check = [(u, p, t) for u, p, t in accounts
                if u not in all_results or all_results[u].get("status") != "COMPLETED"]
    skipped = len(accounts) - len(to_check)
    if skipped:
        print(f"Skipping {skipped} already-COMPLETED")

    total = len(to_check)
    print(f"Checking {total} accounts (parallel={PARALLEL})")
    if total == 0:
        print("Nothing to check!")
        return

    done_count = [0]

    def _handle(r):
        all_results[r["username"]] = r
        done_count[0] += 1
        print(f"  Progress: {done_count[0]}/{total}")
        save_results(all_results)

    if PARALLEL <= 1:
        for i, (u, p, t) in enumerate(to_check):
            _handle(check_tenant(u, p, t, i + 1, total))
    else:
        with ThreadPoolExecutor(max_workers=PARALLEL) as ex:
            futs = {ex.submit(check_tenant, u, p, t, i + 1, total): i
                    for i, (u, p, t) in enumerate(to_check)}
            for f in as_completed(futs):
                _handle(f.result())

    # Summary
    vals = list(all_results.values())
    completed = [r for r in vals if r["status"] == "COMPLETED"]
    with_files = [r for r in completed if r.get("matches_found", "0") not in ("0", "")]
    no_files = [r for r in completed if r.get("matches_found", "0") in ("0", "")]

    print(f"\n{'='*60}")
    print(f"RESULTS — {SPN_SHEET}")
    print(f"{'='*60}")
    print(f"  COMPLETED:     {len(completed)}")
    print(f"    Files found: {len(with_files)}")
    print(f"    No files:    {len(no_files)}")
    print(f"  RUNNING:       {sum(1 for r in vals if r['status'] == 'RUNNING')}")
    print(f"  NO_DISCOVERY:  {sum(1 for r in vals if r['status'] == 'NO_DISCOVERY')}")
    print(f"  ERROR:         {sum(1 for r in vals if r['status'] in ('ERROR','CHECK_ERROR','FAILED'))}")
    print(f"  TOTAL:         {len(vals)}")
    print(f"\nSaved: {RESULTS_FILE}")

    if with_files:
        print(f"\nAccounts with files found:")
        for r in with_files:
            print(f"  {r['username']}: {r['matches_found']} files — {r.get('summary','')[:80]}")


if __name__ == "__main__":
    main()
