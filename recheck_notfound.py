#!/usr/bin/env python3
"""
Re-check NOT_FOUND accounts (non-permission issues only).
Handles: blank pages, classic portal dialog, login failures, page timeouts.
Runs one at a time with enhanced interactive handling.
Screenshots saved with _v2 suffix to preserve originals.
"""

import time
import sys
import os
import csv
import json
import openpyxl
from datetime import datetime
from playwright.sync_api import sync_playwright

SPN_FILE = r"C:\certs\spns.xlsx"
OUTPUT_CSV = r"C:\certs\recheck_results.csv"

# Accounts to recheck grouped by batch (non-permission NOT_FOUND only)
RECHECK_ACCOUNTS = {
    "Batch1": [
        "odl_user_2226992",  # Switch to classic portal dialog
        "odl_user_2227021",  # Switch to classic portal dialog
        "odl_user_2227233",  # Switch to classic portal dialog
        "odl_user_2227843",  # Switch to classic portal dialog
    ],
    "Batch2": [
        "odl_user_2227039",  # Login failed
        "odl_user_2227043",  # Blank/empty Posture Agent tab
        "odl_user_2227059",  # Blank/empty Posture Agent tab
        "odl_user_2227242",  # Blank/empty Posture Agent tab
        "odl_user_2227844",  # Blank/empty Posture Agent tab
    ],
    "Batch3": [
        "odl_user_2227091",  # Blank/empty Asset explorer page
        "odl_user_2227093",  # Blank/empty Asset explorer page
        "odl_user_2227126",  # Blank/empty Asset explorer page
        "odl_user_2227129",  # Page loading timeout
        "odl_user_2227131",  # Blank/empty Asset explorer page
    ],
    "Batch4": [
        "odl_user_2227168",  # Blank/empty Asset explorer page
        "odl_user_2227178",  # Blank/empty Asset explorer page
        "odl_user_2227848",  # Blank/empty Asset explorer page
        "odl_user_2227855",  # Blank/empty Asset explorer page
    ],
    "Batch5": [
        "odl_user_2227195",  # Blank/empty Posture Agent tab
        "odl_user_2227196",  # Blank/empty Posture Agent tab
        "odl_user_2227322",  # Blank/empty Posture Agent tab
        "odl_user_2227323",  # Blank/empty Posture Agent tab
        "odl_user_2227324",  # Blank/empty Posture Agent tab
        "odl_user_2227327",  # Blank/empty Posture Agent tab
        "odl_user_2227328",  # Blank/empty Posture Agent tab
        "odl_user_2227332",  # Blank/empty Posture Agent tab
        "odl_user_2227342",  # Blank/empty Posture Agent tab
        "odl_user_2227345",  # Blank/empty Posture Agent tab
        "odl_user_2227414",  # Blank/empty Posture Agent tab
        "odl_user_2227418",  # Blank/empty Posture Agent tab
        "odl_user_2227421",  # Blank/empty Posture Agent tab
        "odl_user_2227495",  # Blank/empty Posture Agent tab
        "odl_user_2227507",  # Blank/empty Posture Agent tab
        "odl_user_2227508",  # Blank/empty Posture Agent tab
        "odl_user_2227510",  # Blank/empty Posture Agent tab
        "odl_user_2227513",  # Blank/empty Posture Agent tab
        "odl_user_2227821",  # Blank/empty Posture Agent tab
    ],
}


def load_credentials(batch, user_prefix):
    """Load credentials for a specific user from spns.xlsx."""
    wb = openpyxl.load_workbook(SPN_FILE, read_only=True, data_only=True)
    ws = wb[batch]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {h: i for i, h in enumerate(headers)}
    for row in ws.iter_rows(min_row=2, values_only=True):
        odluser = str(row[col["odluser"]] or "").strip()
        if odluser.startswith(user_prefix + "@"):
            wb.close()
            return {
                "username": odluser,
                "password": str(row[col["odlpassword"]] or "").strip(),
                "tenant_id": str(row[col["TenantId"]] or "").strip(),
            }
    wb.close()
    return None


def dismiss_all_dialogs(page):
    """Aggressively dismiss all dialogs, popups, and overlays."""
    # Classic portal switch dialog
    for sel in [
        'button:has-text("Switch to the new Microsoft Purview portal")',
        'button:has-text("Switch")',
        'a:has-text("Switch to the new Microsoft Purview portal")',
        'a:has-text("Try the new")',
        'button:has-text("Try now")',
        'a:has-text("Try now")',
    ]:
        try:
            btn = page.locator(sel)
            if btn.count() > 0 and btn.first.is_visible():
                btn.first.click(timeout=3000)
                print(f"    -> Clicked: {sel}")
                time.sleep(3)
        except Exception:
            pass

    # Standard close/dismiss buttons
    for sel in [
        'button[aria-label="Close"]', 'button[aria-label="Dismiss"]',
        'button:has-text("Close")', 'button:has-text("Dismiss")',
        'button:has-text("Got it")', 'button:has-text("Skip")',
        'button:has-text("OK")', 'button:has-text("No thanks")',
        'button:has-text("Maybe later")', 'button:has-text("Not now")',
        'button:has-text("Don\'t show again")',
        'div[role="dialog"] button[aria-label="Close"]',
    ]:
        try:
            btn = page.locator(sel)
            if btn.count() > 0 and btn.first.is_visible():
                btn.first.click(timeout=2000)
                print(f"    -> Dismissed: {sel}")
                time.sleep(1)
        except Exception:
            pass

    # Copilot panel
    try:
        cp = page.locator('button[aria-label="Close Security Copilot"], button[aria-label="Close Copilot"], button[aria-label="Close panel"]')
        if cp.count() > 0 and cp.first.is_visible():
            cp.first.click(timeout=3000, force=True)
            print("    -> Closed Copilot panel")
            time.sleep(1)
    except Exception:
        pass

    # Remove backdrop overlays via JS
    try:
        page.evaluate('''() => {
            document.querySelectorAll(".fui-DialogSurface__backdrop, .ms-Overlay, .ms-Modal-backdrop").forEach(e => e.remove());
            document.querySelectorAll('[class*="backdrop"], [class*="overlay"]').forEach(e => {
                if (e.style.position === "fixed" || e.style.position === "absolute") e.remove();
            });
        }''')
    except Exception:
        pass


def handle_classic_portal(page):
    """Handle 'Switch to classic portal' or compliance portal redirects."""
    try:
        url = page.url
        if "compliance.microsoft.com" in url or "classic" in url:
            print("    -> Detected classic portal URL, navigating back to new portal")
            return True
    except Exception:
        pass

    # Check for classic portal dialog text
    try:
        classic_text = page.locator('text=/classic|old portal|compliance portal/i')
        if classic_text.count() > 0:
            # Try to find "switch" or "try new" button
            for sel in [
                'button:has-text("Switch")', 'a:has-text("Switch")',
                'button:has-text("Try")', 'a:has-text("Try")',
                'button:has-text("new portal")', 'a:has-text("new portal")',
            ]:
                try:
                    btn = page.locator(sel)
                    if btn.count() > 0 and btn.first.is_visible():
                        btn.first.click(timeout=3000)
                        print(f"    -> Switched via: {sel}")
                        time.sleep(5)
                        return True
                except Exception:
                    pass
    except Exception:
        pass
    return False


def handle_get_started(page):
    """Handle first-time setup dialogs."""
    try:
        gs = page.locator('button:has-text("Get started")')
        if gs.count() > 0 and gs.first.is_visible():
            gs.first.click(timeout=5000)
            print("    -> Clicked 'Get started'")
            time.sleep(10)
            # Look for Start setup
            try:
                ss = page.locator('button:has-text("Start setup")')
                if ss.count() > 0 and ss.first.is_visible():
                    ss.first.click(timeout=5000)
                    print("    -> Clicked 'Start setup'")
                    time.sleep(12)
            except Exception:
                pass
            # Close any resulting dialog
            try:
                cb = page.locator('button:has-text("Close")')
                if cb.count() > 0 and cb.first.is_visible():
                    cb.first.click(timeout=5000)
                    time.sleep(3)
            except Exception:
                pass
            return True
    except Exception:
        pass
    return False


def detect_page_state(page):
    """Detect what's on the page and return a description."""
    checks = {}

    # Check URL
    try:
        checks["url"] = page.url
    except Exception:
        checks["url"] = "unknown"

    # Check for View insights (COMPLETED)
    try:
        checks["view_insights"] = page.locator('button:has-text("View insights")').count() > 0
    except Exception:
        checks["view_insights"] = False

    # Check for Pick up where you left off (RUNNING)
    try:
        checks["pickup"] = page.locator('text=/Pick up where you left off/i').count() > 0
    except Exception:
        checks["pickup"] = False

    # Check for running indicators
    try:
        checks["running"] = page.locator('text=/[Rr]unning|[Ii]n progress|[Pp]rocessing|[Ss]top generating|[Jj]ob estimation/').count() > 0
    except Exception:
        checks["running"] = False

    # Check for Add data sources
    try:
        checks["add_sources"] = page.locator('button:has-text("Add data sources"), button:has-text("Add data source")').count() > 0
    except Exception:
        checks["add_sources"] = False

    # Check for Posture Agent tab
    try:
        checks["pa_tab"] = page.locator('button:has-text("Posture Agent"), button:has-text("Posture agent")').count() > 0
    except Exception:
        checks["pa_tab"] = False

    # Check for permission errors
    try:
        checks["no_permission"] = page.locator('text=/don\'t have permission|not authorized|access denied|insufficient/i').count() > 0
    except Exception:
        checks["no_permission"] = False

    # Check for classic portal dialog
    try:
        checks["classic_dialog"] = page.locator('text=/classic|Switch to|compliance portal/i').count() > 0
    except Exception:
        checks["classic_dialog"] = False

    # Check for loading spinners
    try:
        checks["loading"] = page.locator('.ms-Shimmer-container, .ms-Spinner, [class*="spinner"], [class*="loading"], [class*="Shimmer"]').count() > 0
    except Exception:
        checks["loading"] = False

    # Check for login page
    try:
        checks["login_page"] = page.locator('input[type="email"], input[name="loginfmt"]').count() > 0
    except Exception:
        checks["login_page"] = False

    # Check for error banners
    try:
        checks["error_banner"] = page.locator('text=/something went wrong|error occurred|try again/i').count() > 0
    except Exception:
        checks["error_banner"] = False

    # Check for Discovery agent card
    try:
        checks["discovery_card"] = page.locator('text=/Discovery agent/i, text=/Sensitive/i').count() > 0
    except Exception:
        checks["discovery_card"] = False

    return checks


def determine_status(checks):
    """Determine the actual status from page state checks."""
    if checks.get("view_insights"):
        return "COMPLETED", "View insights button found"
    if checks.get("running"):
        return "RUNNING", "Job running/in progress"
    if checks.get("pickup"):
        return "RUNNING", "Pick up where you left off present"
    if checks.get("add_sources") or checks.get("discovery_card"):
        return "RUNNING", "Data sources/discovery card visible"
    if checks.get("no_permission"):
        return "NO_PERMISSION", "Insufficient permissions"
    if checks.get("login_page"):
        return "LOGIN_FAILED", "Stuck on login page"
    if checks.get("loading"):
        return "LOADING", "Page still loading"
    if checks.get("classic_dialog"):
        return "CLASSIC_DIALOG", "Classic portal dialog blocking"
    if checks.get("error_banner"):
        return "ERROR", "Error banner on page"
    return "NOT_FOUND", "No DSPM content detected"


def recheck_account(batch, user_prefix, index, total):
    """Re-check a single NOT_FOUND account with enhanced handling."""
    creds = load_credentials(batch, user_prefix)
    if not creds:
        print(f"  [SKIP] Could not find credentials for {user_prefix} in {batch}")
        return {"user": user_prefix, "batch": batch, "status": "SKIP", "remark": "Credentials not found"}

    username = creds["username"]
    password = creds["password"]
    tenant_id = creds["tenant_id"]

    print(f"\n{'='*70}")
    print(f"[{index}/{total}] {user_prefix} ({batch}) - TID: {tenant_id}")
    print(f"{'='*70}")

    screenshots_dir = rf"C:\certs\screenshots_{batch}"
    os.makedirs(screenshots_dir, exist_ok=True)

    result = {
        "user": user_prefix,
        "batch": batch,
        "tenant_id": tenant_id,
        "status": "ERROR",
        "remark": "",
        "screenshot": "",
        "checked_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }

    browser = None
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=False)
            context = browser.new_context(viewport={"width": 1920, "height": 1080})
            page = context.new_page()

            # === LOGIN ===
            print("  [1/5] Logging in...")
            page.goto("https://purview.microsoft.com", timeout=120000)
            time.sleep(3)
            try:
                page.wait_for_load_state("load", timeout=30000)
            except Exception:
                pass

            # Fill email
            try:
                page.fill('input[type="email"]', username, timeout=15000)
                page.click('input[type="submit"]')
                time.sleep(4)
            except Exception as e:
                print(f"  [ERROR] Email fill failed: {e}")
                result["remark"] = f"Login error: {e}"
                return result

            # Switch to password if needed
            try:
                pw_link = page.locator('a:has-text("Use your password instead"), a:has-text("password"), #idA_PWD_SwitchToPassword')
                pw_link.first.click(timeout=8000)
                time.sleep(3)
            except Exception:
                pass

            # Fill password
            try:
                page.fill('input[type="password"]', password)
                page.click('input[type="submit"]')
                time.sleep(5)
            except Exception as e:
                print(f"  [ERROR] Password fill failed: {e}")
                result["remark"] = f"Password error: {e}"
                return result

            # Stay signed in
            try:
                yes_btn = page.locator('input[type="submit"][value="Yes"], button:has-text("Yes")')
                yes_btn.first.click(timeout=8000)
                time.sleep(3)
            except Exception:
                pass

            try:
                page.wait_for_load_state("load", timeout=30000)
            except Exception:
                pass
            time.sleep(3)

            # Check if still on login page
            state = detect_page_state(page)
            if state.get("login_page"):
                print("  [WARN] Still on login page, trying password again...")
                try:
                    page.fill('input[type="password"]', password)
                    page.click('input[type="submit"]')
                    time.sleep(8)
                except Exception:
                    pass
                try:
                    yes_btn = page.locator('input[type="submit"][value="Yes"], button:has-text("Yes")')
                    yes_btn.first.click(timeout=8000)
                    time.sleep(3)
                except Exception:
                    pass

            print("  [2/5] Login done, dismissing popups...")
            dismiss_all_dialogs(page)
            handle_get_started(page)
            dismiss_all_dialogs(page)

            # === NAVIGATE TO DSPM ===
            print("  [3/5] Navigating to DSPM Asset Explorer...")
            dspm_url = f"https://purview.microsoft.com/dspm/assetexplorer?tid={tenant_id}"
            page.goto(dspm_url, wait_until="domcontentloaded", timeout=120000)
            try:
                page.wait_for_load_state("load", timeout=30000)
            except Exception:
                pass
            time.sleep(10)

            # Dismiss dialogs again after navigation
            dismiss_all_dialogs(page)
            handle_classic_portal(page)
            dismiss_all_dialogs(page)
            handle_get_started(page)
            dismiss_all_dialogs(page)

            # If redirected to classic/compliance portal, re-navigate
            current_url = page.url
            if "compliance.microsoft.com" in current_url:
                print("  [ACTION] Redirected to compliance portal, re-navigating to new portal...")
                page.goto(dspm_url, wait_until="domcontentloaded", timeout=120000)
                time.sleep(10)
                dismiss_all_dialogs(page)

            # === CLICK POSTURE AGENT TAB ===
            print("  [4/5] Looking for Posture Agent tab...")
            pa_clicked = False
            for wait_round in range(24):  # up to 120s
                dismiss_all_dialogs(page)
                handle_classic_portal(page)

                try:
                    pa = page.locator('button:has-text("Posture Agent"), button:has-text("Posture agent")')
                    if pa.count() > 0 and pa.first.is_visible():
                        pa.first.click(timeout=5000, force=True)
                        print(f"    -> Clicked Posture Agent tab (after {wait_round*5}s)")
                        pa_clicked = True
                        break
                except Exception:
                    pass

                # Check if page has content already
                state = detect_page_state(page)
                if state.get("view_insights") or state.get("pickup") or state.get("running"):
                    print(f"    -> Content already visible (no tab click needed)")
                    break
                if state.get("no_permission"):
                    print(f"    -> Permission issue detected")
                    break

                if state.get("loading") and wait_round % 4 == 0:
                    print(f"    -> Page loading... ({wait_round*5}s)")

                # If blank for too long, try refresh once at 60s
                if wait_round == 12:
                    print("    -> Page blank for 60s, refreshing...")
                    try:
                        page.reload(wait_until="domcontentloaded", timeout=60000)
                        time.sleep(10)
                        dismiss_all_dialogs(page)
                        handle_classic_portal(page)
                        dismiss_all_dialogs(page)
                    except Exception:
                        pass

                time.sleep(5)

            if pa_clicked:
                time.sleep(5)

            # === WAIT FOR CONTENT ===
            print("  [5/5] Checking page content...")
            dismiss_all_dialogs(page)

            # Wait up to 60s for content after clicking PA tab
            final_status = "NOT_FOUND"
            final_remark = ""
            for content_wait in range(12):
                dismiss_all_dialogs(page)
                state = detect_page_state(page)
                status, remark = determine_status(state)

                if status in ("COMPLETED", "RUNNING", "NO_PERMISSION"):
                    final_status = status
                    final_remark = remark
                    print(f"    -> Status: {status} - {remark}")
                    break

                if state.get("loading") and content_wait % 4 == 0:
                    print(f"    -> Content loading... ({content_wait*5}s)")

                time.sleep(5)
            else:
                # Final check
                state = detect_page_state(page)
                status, remark = determine_status(state)
                final_status = status
                final_remark = remark
                print(f"    -> Final status: {status} - {remark}")

            # Extra info for COMPLETED
            if final_status == "COMPLETED":
                try:
                    items = page.locator('text=/found \\d+ items/')
                    if items.count() > 0:
                        final_remark += f" | {items.first.inner_text()}"
                except Exception:
                    pass
                try:
                    date_el = page.locator('text=/\\d+\\/\\d+\\/\\d+/')
                    if date_el.count() > 0:
                        final_remark += f" | Date: {date_el.first.inner_text()}"
                except Exception:
                    pass

            result["status"] = final_status
            result["remark"] = final_remark

            # === SCREENSHOT ===
            screenshot_path = os.path.join(screenshots_dir, f"{user_prefix}_{final_status}_v2.png")
            try:
                page.screenshot(path=screenshot_path, full_page=True)
                result["screenshot"] = screenshot_path
                print(f"  Screenshot: {screenshot_path}")
            except Exception as e:
                print(f"  [WARN] Screenshot failed: {e}")

    except Exception as e:
        print(f"  [ERROR] {e}")
        result["status"] = "ERROR"
        result["remark"] = str(e)[:200]
    finally:
        try:
            if browser:
                browser.close()
        except Exception:
            pass

    print(f"  RESULT: {result['status']} - {result['remark']}")
    return result


def save_results(results):
    """Save recheck results to CSV."""
    fieldnames = ["user", "batch", "tenant_id", "status", "remark", "screenshot", "checked_at"]
    with open(OUTPUT_CSV, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(results)
    print(f"\nResults saved to: {OUTPUT_CSV} ({len(results)} entries)")


def main():
    batch_filter = sys.argv[1] if len(sys.argv) > 1 else None

    # Build flat list of accounts to check
    accounts = []
    for batch, users in RECHECK_ACCOUNTS.items():
        if batch_filter and batch != batch_filter:
            continue
        for user in users:
            accounts.append((batch, user))

    total = len(accounts)
    print(f"Re-checking {total} NOT_FOUND accounts (non-permission)")
    if batch_filter:
        print(f"  Filter: {batch_filter} only")

    results = []
    for i, (batch, user) in enumerate(accounts, 1):
        r = recheck_account(batch, user, i, total)
        results.append(r)
        save_results(results)  # Save after each to avoid losing progress

    # Final summary
    print(f"\n{'='*70}")
    print("RECHECK SUMMARY")
    print(f"{'='*70}")
    from collections import Counter
    status_counts = Counter(r["status"] for r in results)
    for status, count in sorted(status_counts.items()):
        print(f"  {status}: {count}")
    print(f"  TOTAL: {total}")


if __name__ == "__main__":
    main()
