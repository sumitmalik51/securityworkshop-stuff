import time
import openpyxl
from playwright.sync_api import sync_playwright

wb = openpyxl.load_workbook('C:/certs/spns.xlsx', read_only=True)
ws = wb['Batch1']
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
row = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
user = str(row[headers.index('odluser')]).strip()
pwd = str(row[headers.index('odlpassword')]).strip()
tid = str(row[headers.index('TenantId')]).strip()
wb.close()
print(f"User: {user}, TID: {tid}")

with sync_playwright() as p:
    browser = p.chromium.launch(headless=False)
    page = browser.new_context().new_page()
    page.goto('https://purview.microsoft.com')
    time.sleep(3)
    page.fill('input[type="email"]', user)
    page.click('input[type="submit"]')
    time.sleep(4)
    try:
        page.locator('a:has-text("Use your password instead"), #idA_PWD_SwitchToPassword').first.click(timeout=8000)
        time.sleep(3)
    except:
        pass
    page.wait_for_selector('input[type="password"]', timeout=10000)
    page.fill('input[type="password"]', pwd)
    page.click('input[type="submit"]')
    time.sleep(4)
    try:
        page.locator('input[type="submit"][value="Yes"]').first.click(timeout=15000)
    except:
        pass
    time.sleep(8)

    # Close popups
    for sel in ['button:has-text("Get started")', 'button[aria-label="Close"]', 'button:has-text("Close")']:
        try:
            b = page.locator(sel)
            if b.count() > 0 and b.first.is_visible():
                b.first.click(timeout=2000)
                time.sleep(1)
        except:
            pass

    # Navigate to DSPM Asset Explorer
    page.goto(f'https://purview.microsoft.com/dspm/assetexplorer?tid={tid}', wait_until='domcontentloaded', timeout=60000)
    time.sleep(10)

    for sel in ['button[aria-label="Close"]', 'button:has-text("Close")']:
        try:
            b = page.locator(sel)
            if b.count() > 0 and b.first.is_visible():
                b.first.click(timeout=2000)
                time.sleep(1)
        except:
            pass

    page.screenshot(path="C:/certs/debug_1_asset_explorer.png", full_page=True)
    print("Screenshot 1: Asset Explorer page saved")

    # Click Posture agent
    page.get_by_text('Posture agent', exact=False).first.click(timeout=10000)
    time.sleep(5)
    page.screenshot(path="C:/certs/debug_2_posture_agent.png", full_page=True)
    print("Screenshot 2: Posture agent page saved")

    # Click Add data
    page.locator('button:has-text("Add data")').first.click(timeout=10000)
    time.sleep(5)
    page.screenshot(path="C:/certs/debug_3_add_data.png", full_page=True)
    print("Screenshot 3: After Add data click saved")

    # Search for user in the correct search box (placeholder='Search for people, groups, or sites')
    search_box = page.locator('input[placeholder="Search for people, groups, or sites"]')
    search_box.fill(user.split('@')[0])  # e.g. odl_user_2226954
    time.sleep(1)
    search_box.press('Enter')
    time.sleep(5)

    page.screenshot(path="C:/certs/debug_4_search_results.png", full_page=True)
    print("Screenshot 4: After search saved")

    # Dump all visible elements in results area
    print("\n=== SEARCH RESULTS AREA ===")
    print(page.locator('body').inner_text()[:4000])

    print("\n=== ALL CHECKBOXES AFTER SEARCH ===")
    cbs = page.locator('input[type="checkbox"], [role="checkbox"]')
    for i in range(min(cbs.count(), 20)):
        try:
            cb = cbs.nth(i)
            label = cb.evaluate('el => { let l = el.closest("label") || el.parentElement; return l ? l.textContent.trim().substring(0,100) : "no-label" }')
            checked = cb.evaluate('el => el.checked || el.getAttribute("aria-checked")')
            vis = cb.is_visible()
            print(f"  cb[{i}]: label='{label}', checked={checked}, visible={vis}")
        except:
            pass

    print("\n=== ALL BUTTONS AFTER SEARCH ===")
    buttons = page.locator('button')
    for i in range(min(buttons.count(), 40)):
        try:
            btn = buttons.nth(i)
            if btn.is_visible():
                txt = btn.text_content().strip()[:100]
                if txt:
                    print(f"  btn[{i}]: '{txt}'")
        except:
            pass

    # Try clicking on a search result if any
    print("\n=== CLICKABLE ITEMS (divs with user text) ===")
    items = page.locator(f'text={user.split("@")[0]}')
    print(f"  Found {items.count()} items with user text")
    for i in range(items.count()):
        try:
            el = items.nth(i)
            tag = el.evaluate('el => el.tagName')
            txt = el.text_content().strip()[:120]
            print(f"  [{i}]: <{tag}> '{txt}'")
        except:
            pass

    # Try selecting the first result
    if items.count() > 0:
        print("\nClicking first user result...")
        items.first.click()
        time.sleep(3)
        page.screenshot(path="C:/certs/debug_5_after_select.png", full_page=True)
        print("Screenshot 5: After selecting user saved")

        print("\n=== CHECKBOXES AFTER SELECT ===")
        cbs = page.locator('input[type="checkbox"], [role="checkbox"]')
        for i in range(min(cbs.count(), 20)):
            try:
                cb = cbs.nth(i)
                label = cb.evaluate('el => { let l = el.closest("label") || el.parentElement; return l ? l.textContent.trim().substring(0,100) : "no-label" }')
                checked = cb.evaluate('el => el.checked || el.getAttribute("aria-checked")')
                vis = cb.is_visible()
                print(f"  cb[{i}]: label='{label}', checked={checked}, visible={vis}")
            except:
                pass

    # Keep open so user can see
    time.sleep(600)
    browser.close()
