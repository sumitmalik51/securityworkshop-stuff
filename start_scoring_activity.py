import time
import sys
import os
import csv
import openpyxl
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from playwright.sync_api import sync_playwright

SPN_FILE = r"C:\certs\spns.xlsx"
SPN_SHEET = sys.argv[1] if len(sys.argv) > 1 else "Batch1"
START_INDEX = int(sys.argv[2]) if len(sys.argv) > 2 else 0
MAX_COUNT = int(sys.argv[3]) if len(sys.argv) > 3 else 0
PARALLEL = int(sys.argv[4]) if len(sys.argv) > 4 else 1

RESULTS_FILE = rf"C:\certs\scoring_activity_{SPN_SHEET}.csv"


def load_accounts():
    accounts = []
    wb = openpyxl.load_workbook(SPN_FILE, read_only=True)
    ws = wb[SPN_SHEET]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    user_col = headers.index("odluser")
    pwd_col = headers.index("odlpassword")
    tid_col = headers.index("TenantId")
    for row in ws.iter_rows(min_row=2, values_only=True):
        username = row[user_col]
        password = row[pwd_col]
        tenant_id = row[tid_col]
        if username and password and tenant_id:
            accounts.append((str(username).strip(), str(password).strip(), str(tenant_id).strip()))
    wb.close()
    return accounts


def close_popups(page):
    """Close welcome/tour/copilot popups."""
    # Handle "Welcome to new Purview portal" - click Get started
    try:
        btn = page.locator('button:has-text("Get started")')
        if btn.count() > 0 and btn.first.is_visible():
            btn.first.click(timeout=3000)
            time.sleep(3)
    except Exception:
        pass

    for _ in range(5):
        closed = False
        for sel in [
            'button[aria-label="Close"]', 'button[aria-label="Dismiss"]',
            'button:has-text("Close")', 'button:has-text("Dismiss")',
            'button:has-text("Got it")', 'button:has-text("Skip")',
            'button:has-text("OK")', 'button:has-text("No thanks")',
            'button:has-text("Maybe later")', 'button:has-text("Not now")',
            'div[role="dialog"] button[aria-label="Close"]',
        ]:
            try:
                btn = page.locator(sel)
                if btn.count() > 0 and btn.first.is_visible():
                    btn.first.click(timeout=2000)
                    closed = True
                    time.sleep(1)
            except Exception:
                pass
        if not closed:
            break

    try:
        page.evaluate('document.querySelectorAll(".fui-DialogSurface__backdrop").forEach(e => e.remove())')
    except Exception:
        pass


def process_tenant(username, password, tenant_id, index, total):
    """Start scoring activity for a single tenant."""
    print(f"\n=== [{index}/{total}] {username} ===")
    # Extract short username like "odl_user_2226972"
    short_user = username.split("@")[0]
    status = "ERROR"
    error_msg = ""

    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=False)
            context = browser.new_context()
            page = context.new_page()

            # Login
            page.goto("https://purview.microsoft.com")
            time.sleep(3)

            page.fill('input[type="email"]', username)
            page.click('input[type="submit"]')
            time.sleep(4)

            try:
                pw_link = page.locator('a:has-text("Use your password instead"), a:has-text("password"), #idA_PWD_SwitchToPassword')
                pw_link.first.click(timeout=8000)
                time.sleep(3)
            except Exception:
                pass

            page.fill('input[type="password"]', password)
            page.click('input[type="submit"]')
            time.sleep(5)

            # Stay signed in
            try:
                yes_btn = page.locator('input[type="submit"][value="Yes"], button:has-text("Yes")')
                yes_btn.first.click(timeout=8000)
                time.sleep(3)
            except Exception:
                pass

            # Handle Welcome to Purview portal after login
            close_popups(page)
            time.sleep(3)

            # Navigate to IRM Policies page
            irm_url = f"https://purview.microsoft.com/insiderriskmgmt/policiespage?tid={tenant_id}"
            print(f"  Navigating to: {irm_url}")
            page.goto(irm_url, wait_until="domcontentloaded", timeout=60000)
            time.sleep(10)

            close_popups(page)

            # Wait for policies to load (shimmers to disappear, DetailsRow to appear)
            print("  Waiting for policies to load...")
            for _wait in range(15):  # up to 30s
                row_count = page.locator('.ms-DetailsRow').count()
                if row_count > 0:
                    print(f"  Policies loaded ({row_count} rows)")
                    break
                close_popups(page)
                time.sleep(2)
            else:
                print("  [WARN] Policies may not have loaded (no .ms-DetailsRow found)")
            time.sleep(2)

            # Step 1: Select all policies via the header checkbox (select all)
            print("  Selecting all policies...")
            try:
                # The header "Select all" checkbox: div[role="checkbox"] with data-selection-toggle
                header_cb = page.locator('div[role="checkbox"][data-selection-toggle="true"]')
                if header_cb.count() > 0:
                    header_cb.first.click(timeout=5000, force=True)
                    print("  Clicked header select-all checkbox")
                    time.sleep(2)
                else:
                    # Fallback: click individual row checkboxes
                    row_checks = page.locator('.ms-DetailsRow .ms-DetailsRow-check')
                    count = row_checks.count()
                    print(f"  No header checkbox, clicking {count} row checkboxes")
                    for i in range(count):
                        try:
                            row_checks.nth(i).click(timeout=3000, force=True)
                            time.sleep(0.5)
                        except Exception:
                            pass
            except Exception as e:
                print(f"  [WARN] Checkbox selection issue: {e}")

            time.sleep(2)

            # Step 2: Click "Start scoring activity for users"
            print("  Clicking 'Start scoring activity for users'...")
            try:
                score_btn = page.locator('button:has-text("Start scoring activity for users"), button:has-text("Start scoring activity")')
                score_btn.first.click(timeout=10000, force=True)
                print("  Clicked Start scoring activity for users")
                time.sleep(5)
            except Exception as e:
                print(f"  [ERROR] Could not click Start scoring activity: {e}")
                # Try via command bar
                try:
                    page.evaluate('''
                        const btns = document.querySelectorAll('button');
                        for (const b of btns) {
                            if (b.textContent.includes('Start scoring activity')) {
                                b.click();
                                break;
                            }
                        }
                    ''')
                    print("  Clicked via JS fallback")
                    time.sleep(5)
                except Exception:
                    error_msg = "Could not click Start scoring activity"
                    raise

            # Step 3: Fill "Reason for scoring activity" with "workshop"
            print("  Filling reason: workshop...")
            try:
                # The reason field - try multiple selectors
                reason_input = page.locator('input[aria-label*="Reason"], input[placeholder*="reason"]')
                if reason_input.count() == 0:
                    # Try any text input in the dialog
                    reason_input = page.locator('.ms-Panel input[type="text"], [role="dialog"] input[type="text"], div.ms-TextField input')
                if reason_input.count() > 0:
                    reason_input.first.fill("workshop", timeout=10000)
                    print("  Filled reason: workshop")
                else:
                    # Click label then type
                    page.locator('text=Reason for scoring activity').click()
                    time.sleep(1)
                    page.keyboard.press("Tab")
                    time.sleep(0.5)
                    page.keyboard.type("workshop")
                    print("  Typed reason via keyboard")
                time.sleep(2)
            except Exception as e:
                print(f"  [WARN] Reason input issue: {e}")
                try:
                    page.locator('text=Reason for scoring activity').click()
                    time.sleep(1)
                    page.keyboard.press("Tab")
                    time.sleep(0.5)
                    page.keyboard.type("workshop")
                    print("  Typed reason via keyboard fallback")
                except Exception:
                    pass

            # Step 4: Search for ODL user in "Score activity for these users"
            print(f"  Searching for '{short_user}'...")
            try:
                search_input = page.locator('input[placeholder*="Search by name or user principal name"], input[placeholder*="Search by name"]')
                if search_input.count() == 0:
                    # Try broader search
                    search_input = page.locator('.ms-Panel input[type="text"]:nth-of-type(2), [role="dialog"] input[type="text"]:nth-of-type(2)')
                if search_input.count() == 0:
                    # Get all text inputs, the search is typically the second one
                    all_inputs = page.locator('input[type="text"]')
                    if all_inputs.count() >= 2:
                        search_input = all_inputs.nth(1)
                    else:
                        search_input = all_inputs.last
                    search_input.fill(short_user, timeout=10000)
                else:
                    search_input.first.fill(short_user, timeout=10000)
                print(f"  Filled search: {short_user}")
                time.sleep(3)
                page.keyboard.press("Enter")
                time.sleep(8)
            except Exception as e:
                print(f"  [WARN] User search issue: {e}")

            # Step 5: Select the user from results
            print("  Selecting user from results...")
            user_num = short_user.replace("odl_user_", "")
            odl_display = f"ODL_User {user_num}"
            try:
                # Try clicking on suggestion items
                suggestion = page.locator(f'.ms-Suggestions-item, [role="option"], button:has-text("{odl_display}"), button:has-text("{short_user}")')
                if suggestion.count() > 0:
                    suggestion.first.click(timeout=10000)
                    print(f"  Selected {odl_display}")
                else:
                    # Try any element containing the user
                    user_el = page.locator(f'text=/{user_num}/')
                    if user_el.count() > 0:
                        user_el.last.click(timeout=5000)
                        print(f"  Selected via text match")
                    else:
                        # JS fallback
                        page.evaluate(f'''
                            const items = document.querySelectorAll('.ms-Suggestions-item, [role="option"], .ms-PeoplePicker-result');
                            for (const item of items) {{
                                if (item.textContent.includes("{user_num}") || item.textContent.toLowerCase().includes("{short_user.lower()}")) {{
                                    item.click();
                                    break;
                                }}
                            }}
                        ''')
                        print("  Selected via JS fallback")
                time.sleep(3)
            except Exception as e:
                print(f"  [WARN] User selection issue: {e}")

            # Step 6: Click "Start scoring activity" button at the bottom
            print("  Clicking 'Start scoring activity' button...")
            time.sleep(2)
            try:
                # The bottom button is specifically "Start scoring activity" (not "for users")
                start_btn = page.locator('button:has-text("Start scoring activity")')
                # Get the last one (bottom of dialog)
                count = start_btn.count()
                if count > 0:
                    start_btn.last.click(timeout=10000, force=True)
                    print("  Clicked Start scoring activity!")
                    time.sleep(5)
                    status = "SUCCESS"
                else:
                    print("  [ERROR] Start scoring activity button not found")
                    error_msg = "Start scoring activity button not found"
            except Exception as e:
                print(f"  [ERROR] Could not click Start scoring activity: {e}")
                error_msg = str(e)[:200]

            # Wait and close
            time.sleep(5)
            browser.close()

    except Exception as e:
        print(f"  ERROR: {e}")
        error_msg = str(e)[:200]

    print(f"  RESULT: {status}")
    return {
        "username": username,
        "tenant_id": tenant_id,
        "status": status,
        "error": error_msg,
        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


def save_results(results):
    fieldnames = ["username", "tenant_id", "status", "error", "timestamp"]
    with open(RESULTS_FILE, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(results)


def main():
    accounts = load_accounts()
    print(f"Loaded {len(accounts)} accounts from sheet '{SPN_SHEET}'")

    if START_INDEX > 0:
        accounts = accounts[START_INDEX:]
    if MAX_COUNT > 0:
        accounts = accounts[:MAX_COUNT]

    total = len(accounts)
    print(f"Processing {total} accounts (start={START_INDEX}, max={MAX_COUNT}, parallel={PARALLEL})")

    results = []
    success_count = 0
    fail_count = 0

    if PARALLEL <= 1:
        for i, (user, pwd, tid) in enumerate(accounts):
            r = process_tenant(user, pwd, tid, i + 1 + START_INDEX, total)
            results.append(r)
            if r["status"] == "SUCCESS":
                success_count += 1
            else:
                fail_count += 1
            print(f"  Progress: {i+1}/{total} (Success={success_count}, Failed={fail_count})")
            save_results(results)
    else:
        futures = {}
        with ThreadPoolExecutor(max_workers=PARALLEL) as executor:
            for i, (user, pwd, tid) in enumerate(accounts):
                f = executor.submit(process_tenant, user, pwd, tid, i + 1 + START_INDEX, total)
                futures[f] = i

            for f in as_completed(futures):
                r = f.result()
                results.append(r)
                if r["status"] == "SUCCESS":
                    success_count += 1
                else:
                    fail_count += 1
                done = success_count + fail_count
                print(f"  Progress: {done}/{total} (Success={success_count}, Failed={fail_count})")
                save_results(results)

    print(f"\n{'='*60}")
    print(f"FINAL SUMMARY - {SPN_SHEET}")
    print(f"{'='*60}")
    print(f"  SUCCESS:  {success_count}")
    print(f"  FAILED:   {fail_count}")
    print(f"  TOTAL:    {len(results)}")
    print(f"\nResults saved to: {RESULTS_FILE}")


if __name__ == "__main__":
    main()
