import time
import sys
import os
import csv
import openpyxl
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from playwright.sync_api import sync_playwright

SPN_FILE = r"C:\certs\spns.xlsx"
SPN_SHEET = sys.argv[1] if len(sys.argv) > 1 else "Batch1"
START_INDEX = int(sys.argv[2]) if len(sys.argv) > 2 else 0
MAX_COUNT = int(sys.argv[3]) if len(sys.argv) > 3 else 0
PARALLEL = int(sys.argv[4]) if len(sys.argv) > 4 else 3

RESULTS_FILE = rf"C:\certs\dspm_status_{SPN_SHEET}.csv"


def load_accounts():
    accounts = []
    wb = openpyxl.load_workbook(SPN_FILE, read_only=True)
    ws = wb[SPN_SHEET]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    user_col = headers.index("odluser")
    pwd_col = headers.index("odlpassword")
    tid_col = headers.index("TenantId")
    for row in ws.iter_rows(min_row=2, values_only=True):
        username = row[user_col]
        password = row[pwd_col]
        tenant_id = row[tid_col]
        if username and password and tenant_id:
            accounts.append((str(username).strip(), str(password).strip(), str(tenant_id).strip()))
    wb.close()
    return accounts


def close_popups(page):
    """Close welcome/tour/copilot popups."""
    try:
        btn = page.locator('button:has-text("Get started")')
        if btn.count() > 0 and btn.first.is_visible():
            btn.first.click(timeout=3000)
            time.sleep(2)
    except Exception:
        pass

    for _ in range(5):
        closed = False
        for sel in [
            'button[aria-label="Close"]', 'button[aria-label="Dismiss"]',
            'button:has-text("Close")', 'button:has-text("Dismiss")',
            'button:has-text("Got it")', 'button:has-text("Skip")',
            'button:has-text("OK")', 'button:has-text("No thanks")',
            'button:has-text("Maybe later")', 'button:has-text("Not now")',
            'div[role="dialog"] button[aria-label="Close"]',
        ]:
            try:
                btn = page.locator(sel)
                if btn.count() > 0 and btn.first.is_visible():
                    btn.first.click(timeout=2000)
                    closed = True
                    time.sleep(1)
            except Exception:
                pass
        if not closed:
            break

    try:
        cp = page.locator('button[aria-label="Close Security Copilot"], button[aria-label="Close Copilot"], button[aria-label="Close panel"]')
        if cp.count() > 0 and cp.first.is_visible():
            cp.first.click(timeout=3000, force=True)
            time.sleep(1)
    except Exception:
        pass

    try:
        page.evaluate('document.querySelectorAll(".fui-DialogSurface__backdrop").forEach(e => e.remove())')
    except Exception:
        pass


def handle_setup_dialogs(page):
    """Handle setup dialogs that may appear for fresh accounts."""
    clicked_get_started = False
    try:
        get_started = page.locator('button:has-text("Get started")')
        if get_started.count() > 0 and get_started.first.is_visible():
            get_started.first.click(timeout=5000)
            clicked_get_started = True
            time.sleep(10)
    except Exception:
        pass
    try:
        start_setup = page.locator('button:has-text("Start setup")')
        if clicked_get_started:
            start_setup.first.wait_for(state="visible", timeout=15000)
        if start_setup.count() > 0 and start_setup.first.is_visible():
            start_setup.first.click(timeout=5000)
            time.sleep(12)
    except Exception:
        pass
    try:
        close_btn = page.locator('button:has-text("Close")')
        if close_btn.count() > 0 and close_btn.first.is_visible():
            close_btn.first.click(timeout=5000)
            time.sleep(3)
    except Exception:
        pass
    if clicked_get_started:
        try:
            page.reload(wait_until="domcontentloaded", timeout=120000)
            time.sleep(15)
        except Exception:
            pass


def check_tenant(username, password, tenant_id, index, total):
    """Check DSPM status for a single tenant. Returns dict with status info."""
    print(f"\n=== [{index}/{total}] {username} ===")
    result = {
        "username": username,
        "tenant_id": tenant_id,
        "status": "ERROR",
        "job_title": "",
        "job_date": "",
        "items_found": "",
        "checked_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }

    for attempt in range(3):
        if attempt > 0:
            print(f"  Retry attempt {attempt+1}/3 for {username}")
            time.sleep(5)
        browser = None
        try:
            with sync_playwright() as p:
                browser = p.chromium.launch(headless=False)
                context = browser.new_context()
                page = context.new_page()

                # Login (retry up to 3 times)
                for _login_nav in range(3):
                    try:
                        page.goto("https://purview.microsoft.com", timeout=120000)
                        break
                    except Exception:
                        if _login_nav == 2:
                            raise
                        time.sleep(10)
                time.sleep(3)

                try:
                    page.wait_for_load_state("load", timeout=30000)
                except Exception:
                    pass

                page.fill('input[type="email"]', username)
                page.click('input[type="submit"]')
                time.sleep(4)

                try:
                    pw_link = page.locator('a:has-text("Use your password instead"), a:has-text("password"), #idA_PWD_SwitchToPassword')
                    pw_link.first.click(timeout=8000)
                    time.sleep(3)
                except Exception:
                    pass

                page.fill('input[type="password"]', password)
                page.click('input[type="submit"]')
                time.sleep(5)

                # Stay signed in
                try:
                    yes_btn = page.locator('input[type="submit"][value="Yes"], button:has-text("Yes")')
                    yes_btn.first.click(timeout=8000)
                    time.sleep(3)
                except Exception:
                    pass

                try:
                    page.wait_for_load_state("load", timeout=30000)
                except Exception:
                    pass
                time.sleep(3)
                close_popups(page)
                url = f"https://purview.microsoft.com/dspm/assetexplorer?tid={tenant_id}"
                print(f"  Navigating to: {url}")
                for _nav_attempt in range(3):
                    try:
                        page.goto(url, wait_until="domcontentloaded", timeout=120000)
                        break
                    except Exception as nav_err:
                        print(f"  [WARN] Navigation attempt {_nav_attempt+1} failed: {nav_err}")
                        if _nav_attempt == 2:
                            raise
                        time.sleep(10)
                    try:
                        page.wait_for_load_state("load", timeout=30000)
                    except Exception:
                        pass
                    time.sleep(15)

                close_popups(page)
                handle_setup_dialogs(page)
                close_popups(page)

                # Click Posture Agent tab - poll up to 180s for slow loading
                pa_clicked = False
                for _pa_wait in range(36):
                    try:
                        pa = page.locator('button:has-text("Posture Agent"), button:has-text("Posture agent")')
                        if pa.count() > 0 and pa.first.is_visible():
                            pa.first.click(timeout=5000, force=True)
                            print("  Clicked Posture Agent tab")
                            pa_clicked = True
                            break
                    except Exception:
                        pass
                    # Check if page is still loading (spinners/shimmers)
                    is_loading = False
                    try:
                        shimmers = page.locator('.ms-Shimmer-container, .ms-Spinner, [class*="spinner"], [class*="loading"]')
                        if shimmers.count() > 0:
                            is_loading = True
                            if _pa_wait % 6 == 0:
                                print(f"  Page still loading (shimmer/spinner detected)... waiting ({_pa_wait*5}s)")
                    except Exception:
                        pass
                    close_popups(page)
                    handle_setup_dialogs(page)
                    time.sleep(5)
                if not pa_clicked:
                    print("  [WARN] Could not click Posture Agent after 180s")

                # Wait for Posture Agent content to fully load (up to 180s)
                print("  Waiting for Posture Agent content to load...")
                content_loaded = False
                for _content_wait in range(36):
                    close_popups(page)
                    handle_setup_dialogs(page)
                    # Check if any recognizable content has loaded
                    has_insights = page.locator('button:has-text("View insights")').count() > 0
                    has_pickup = page.locator('text=/Pick up where you left off/i').count() > 0
                    has_add = page.locator('button:has-text("Add data sources"), button:has-text("Add data source")').count() > 0
                    has_card = page.locator('text=/Discovery agent/i, text=/Sensitive/i').count() > 0
                    if has_insights or has_pickup or has_add or has_card:
                        print(f"  Content loaded (insights={has_insights} pickup={has_pickup} add={has_add} card={has_card})")
                        content_loaded = True
                        break
                    # Check if page is still loading — if so, keep waiting patiently
                    try:
                        shimmers = page.locator('.ms-Shimmer-container, .ms-Spinner, [class*="spinner"], [class*="loading"], [class*="Shimmer"]')
                        if shimmers.count() > 0:
                            if _content_wait % 6 == 0:
                                print(f"  Page still loading (shimmer/spinner detected)... waiting ({_content_wait*5}s)")
                            time.sleep(5)
                            continue
                    except Exception:
                        pass
                    time.sleep(5)
                if not content_loaded:
                    # One final check after extra 30s wait in case of very slow load
                    print("  [WARN] Content not loaded after 180s, doing final 30s wait...")
                    time.sleep(30)
                    close_popups(page)
                    has_insights = page.locator('button:has-text("View insights")').count() > 0
                    has_pickup = page.locator('text=/Pick up where you left off/i').count() > 0
                    has_add = page.locator('button:has-text("Add data sources"), button:has-text("Add data source")').count() > 0
                    if has_insights or has_pickup or has_add:
                        print(f"  Content finally loaded after extra wait (insights={has_insights} pickup={has_pickup} add={has_add})")
                    else:
                        print("  [WARN] Posture Agent content did not load after 210s total")

                close_popups(page)
                handle_setup_dialogs(page)
                close_popups(page)

                # Remove overlays
                try:
                    page.evaluate('document.querySelectorAll(".fui-DialogSurface__backdrop").forEach(e => e.remove())')
                except Exception:
                    pass

                # Now check for "View insights" button - indicates COMPLETED
                time.sleep(3)
                try:
                    view_insights = page.locator('button:has-text("View insights")')
                    if view_insights.count() > 0 and view_insights.first.is_visible():
                        result["status"] = "COMPLETED"
                        print("  STATUS: COMPLETED (View insights available)")

                        # Try to extract job details from the card
                        try:
                            # Get the card text for job title
                            card = page.locator('text=Sensitive Password And Credential Files')
                            if card.count() > 0:
                                result["job_title"] = "Sensitive Password And Credential Files"
                        except Exception:
                            pass

                        # Try to get the items found text
                        try:
                            items_text = page.locator('text=/found \\d+ items/')
                            if items_text.count() > 0:
                                text = items_text.first.inner_text()
                                result["items_found"] = text
                                print(f"  {text}")
                        except Exception:
                            pass

                        # Try to get date
                        try:
                            date_el = page.locator('text=/\\d+\\/\\d+\\/\\d+/')
                            if date_el.count() > 0:
                                result["job_date"] = date_el.first.inner_text()
                        except Exception:
                            pass

                    else:
                        # No View insights - check if there's a running indicator
                        # Look for any card content that says running/in progress
                        running = page.locator('text=/[Rr]unning|[Ii]n progress|[Pp]rocessing|[Ss]top generating|[Jj]ob estimation/')
                        if running.count() > 0:
                            result["status"] = "RUNNING"
                            print("  STATUS: RUNNING")
                        else:
                            # Check if "Pick up where you left off" section exists but no insights
                            pickup = page.locator('text=Pick up where you left off')
                            if pickup.count() > 0:
                                result["status"] = "RUNNING"
                                print("  STATUS: RUNNING (no View insights yet)")
                            else:
                                # No card at all - might not have been set up
                                result["status"] = "NOT_FOUND"
                                print("  STATUS: NOT_FOUND (no job card found)")
                except Exception as e:
                    print(f"  [WARN] Status check error: {e}")
                    result["status"] = "CHECK_ERROR"

        except Exception as e:
            print(f"  ERROR (attempt {attempt+1}): {e}")
            result["status"] = "ERROR"
        finally:
            try:
                if browser:
                    browser.close()
            except Exception:
                pass

        # If we got a real status, stop retrying
        if result["status"] in ("COMPLETED", "RUNNING", "NOT_FOUND"):
            break

    print(f"  RESULT: {result['status']}")
    return result


def load_existing_results():
    """Load existing results from CSV so we can skip already-COMPLETED accounts."""
    existing = {}
    if os.path.exists(RESULTS_FILE):
        try:
            with open(RESULTS_FILE, "r", newline="", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    existing[row["username"]] = row
            print(f"Loaded {len(existing)} existing results from {RESULTS_FILE}")
            completed = sum(1 for r in existing.values() if r["status"] == "COMPLETED")
            print(f"  Already COMPLETED: {completed} (will skip these)")
        except Exception as e:
            print(f"  [WARN] Could not load existing results: {e}")
    return existing


def save_results(all_results):
    """Save all results (merged) to CSV file."""
    fieldnames = ["username", "tenant_id", "status", "job_title", "items_found", "job_date", "checked_at"]
    with open(RESULTS_FILE, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(all_results.values())
    print(f"\nResults saved to: {RESULTS_FILE} ({len(all_results)} entries)")


def main():
    accounts = load_accounts()
    print(f"Loaded {len(accounts)} accounts from sheet '{SPN_SHEET}'")

    # Load existing results to skip already-COMPLETED
    all_results = load_existing_results()

    if START_INDEX > 0:
        accounts = accounts[START_INDEX:]
    if MAX_COUNT > 0:
        accounts = accounts[:MAX_COUNT]

    # Filter out already-COMPLETED accounts
    accounts_to_check = []
    skipped = 0
    for user, pwd, tid in accounts:
        if user in all_results and all_results[user]["status"] == "COMPLETED":
            skipped += 1
        else:
            accounts_to_check.append((user, pwd, tid))
    if skipped > 0:
        print(f"Skipping {skipped} already-COMPLETED accounts")

    total = len(accounts_to_check)
    print(f"Checking {total} accounts (start={START_INDEX}, max={MAX_COUNT}, parallel={PARALLEL})")

    if total == 0:
        print("Nothing to check - all accounts already COMPLETED!")
        save_results(all_results)
        return

    success_count = 0
    fail_count = 0

    def _process_result(r):
        nonlocal success_count, fail_count
        all_results[r["username"]] = r
        if r["status"] == "COMPLETED":
            success_count += 1
        else:
            fail_count += 1
        done = success_count + fail_count
        print(f"  Progress: {done}/{total} (Completed={success_count}, Other={fail_count})")
        save_results(all_results)

    if PARALLEL <= 1:
        for i, (user, pwd, tid) in enumerate(accounts_to_check):
            r = check_tenant(user, pwd, tid, i + 1, total)
            _process_result(r)
    else:
        futures = {}
        with ThreadPoolExecutor(max_workers=PARALLEL) as executor:
            for i, (user, pwd, tid) in enumerate(accounts_to_check):
                f = executor.submit(check_tenant, user, pwd, tid, i + 1, total)
                futures[f] = i

            for f in as_completed(futures):
                r = f.result()
                _process_result(r)

    # Final summary
    print(f"\n{'='*60}")
    print(f"FINAL SUMMARY - {SPN_SHEET}")
    print(f"{'='*60}")
    all_vals = list(all_results.values())
    completed = sum(1 for r in all_vals if r["status"] == "COMPLETED")
    running = sum(1 for r in all_vals if r["status"] == "RUNNING")
    not_found = sum(1 for r in all_vals if r["status"] == "NOT_FOUND")
    errors = sum(1 for r in all_vals if r["status"] in ("ERROR", "CHECK_ERROR"))
    print(f"  COMPLETED:  {completed}")
    print(f"  RUNNING:    {running}")
    print(f"  NOT_FOUND:  {not_found}")
    print(f"  ERRORS:     {errors}")
    print(f"  TOTAL:      {len(all_results)}")
    print(f"\nResults saved to: {RESULTS_FILE}")


if __name__ == "__main__":
    main()
