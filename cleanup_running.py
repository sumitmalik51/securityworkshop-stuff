#!/usr/bin/env python3
"""
Run OneDrive cleanup for RUNNING accounts from dspm_consolidated_status.xlsx.
Uses Graph API (ROPC) to check/cleanup OneDrive root - deletes everything
except 'DSPM-Secret-Scanning-Test' folder.
Adds cleanup remarks back to the Excel.

Usage:
  python cleanup_running.py [start_index] [max_count]
  python cleanup_running.py          # all 190 RUNNING accounts
  python cleanup_running.py 0 5      # first 5 only
"""

import sys
import os
import csv
import json
import requests
import openpyxl
from datetime import datetime

EXCEL_FILE = r"C:\certs\dspm_consolidated_status.xlsx"
RESULTS_FILE = r"C:\certs\cleanup_running_results.csv"
KEEP_FOLDER = "DSPM-Secret-Scanning-Test"
CLIENT_ID = "d3590ed6-52b3-4102-aeff-aad2292ab01c"  # Microsoft Office

START_INDEX = int(sys.argv[1]) if len(sys.argv) > 1 else 0
MAX_COUNT = int(sys.argv[2]) if len(sys.argv) > 2 else 0


def load_running_accounts():
    wb = openpyxl.load_workbook(EXCEL_FILE, read_only=True)
    ws = wb["RUNNING"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    accounts = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        acc = {h: (str(v).strip() if v else "") for h, v in zip(headers, row)}
        accounts.append(acc)
    wb.close()
    return accounts


def get_token(tenant_id, username, password):
    resp = requests.post(
        f"https://login.microsoftonline.com/{tenant_id}/oauth2/v2.0/token",
        data={
            "grant_type": "password",
            "client_id": CLIENT_ID,
            "scope": "https://graph.microsoft.com/.default offline_access",
            "username": username,
            "password": password,
        },
        timeout=30,
    )
    data = resp.json()
    if "access_token" in data:
        return data["access_token"], None
    return None, data.get("error_description", data.get("error", "Unknown auth error"))[:200]


def cleanup_onedrive(token):
    """List OneDrive root, delete everything except KEEP_FOLDER. Return summary."""
    headers = {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}

    # List root items
    items = []
    url = "https://graph.microsoft.com/v1.0/me/drive/root/children?$select=id,name,folder,file,size"
    try:
        while url:
            resp = requests.get(url, headers=headers, timeout=30)
            if resp.status_code != 200:
                return {"status": "LIST_FAILED", "remark": f"HTTP {resp.status_code}: {resp.text[:200]}",
                        "extra_items": 0, "deleted": 0, "items_detail": ""}
            data = resp.json()
            items.extend(data.get("value", []))
            url = data.get("@odata.nextLink")
    except Exception as e:
        return {"status": "LIST_ERROR", "remark": str(e)[:200], "extra_items": 0, "deleted": 0, "items_detail": ""}

    # Find items to delete
    keep_found = False
    extra_items = []
    for item in items:
        if item["name"] == KEEP_FOLDER:
            keep_found = True
        else:
            item_type = "folder" if "folder" in item else "file"
            extra_items.append({"name": item["name"], "type": item_type, "id": item["id"]})

    if not items:
        return {"status": "EMPTY", "remark": "OneDrive root is empty",
                "extra_items": 0, "deleted": 0, "items_detail": ""}

    if not extra_items:
        remark = f"Clean - only '{KEEP_FOLDER}' found" if keep_found else f"Only items: {', '.join(i['name'] for i in items)}"
        return {"status": "CLEAN", "remark": remark, "extra_items": 0, "deleted": 0,
                "items_detail": ", ".join(i["name"] for i in items)}

    # Delete extra items
    deleted = 0
    failed_items = []
    for item in extra_items:
        try:
            del_resp = requests.delete(
                f"https://graph.microsoft.com/v1.0/me/drive/items/{item['id']}",
                headers={"Authorization": f"Bearer {token}"},
                timeout=30,
            )
            if del_resp.status_code in (200, 204):
                deleted += 1
            else:
                failed_items.append(f"{item['name']}(HTTP{del_resp.status_code})")
        except Exception as e:
            failed_items.append(f"{item['name']}(err)")

    items_detail = "; ".join(f"{i['type']}:{i['name']}" for i in extra_items)
    if deleted == len(extra_items):
        status = "CLEANED"
        remark = f"Deleted {deleted} extra items: {items_detail}"
    else:
        status = "PARTIAL"
        remark = f"Deleted {deleted}/{len(extra_items)}. Failed: {', '.join(failed_items)}"

    if not keep_found:
        remark += f" | WARNING: '{KEEP_FOLDER}' not found!"

    return {"status": status, "remark": remark, "extra_items": len(extra_items),
            "deleted": deleted, "items_detail": items_detail}


def process_account(acc, index, total):
    username = acc["ODL User"]
    password = acc["ODL Password"]
    tenant_id = acc["Tenant ID"]
    batch = acc["Batch"]

    print(f"\n{'='*60}")
    print(f"[{index}/{total}] {username} ({batch})")

    result = {
        "batch": batch,
        "username": username,
        "tenant_id": tenant_id,
        "cleanup_status": "ERROR",
        "extra_items": 0,
        "deleted": 0,
        "items_detail": "",
        "remark": "",
        "checked_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }

    # Auth
    token, err = get_token(tenant_id, username, password)
    if not token:
        result["cleanup_status"] = "AUTH_FAILED"
        result["remark"] = err
        print(f"  AUTH_FAILED: {err[:80]}")
        return result

    # Cleanup
    cleanup = cleanup_onedrive(token)
    result["cleanup_status"] = cleanup["status"]
    result["extra_items"] = cleanup["extra_items"]
    result["deleted"] = cleanup["deleted"]
    result["items_detail"] = cleanup["items_detail"]
    result["remark"] = cleanup["remark"]

    print(f"  {cleanup['status']} | extra={cleanup['extra_items']} | deleted={cleanup['deleted']}")
    if cleanup["items_detail"]:
        print(f"  Items: {cleanup['items_detail'][:100]}")

    return result


def save_result_row(result):
    fieldnames = ["batch", "username", "tenant_id", "cleanup_status", "extra_items",
                  "deleted", "items_detail", "remark", "checked_at"]
    file_exists = os.path.exists(RESULTS_FILE) and os.path.getsize(RESULTS_FILE) > 0
    with open(RESULTS_FILE, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        if not file_exists:
            writer.writeheader()
        writer.writerow(result)


def update_excel_remarks(results):
    """Write cleanup remarks back to the RUNNING sheet in consolidated Excel."""
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["RUNNING"]
    headers = [c.value for c in ws[1]]

    remark_col = headers.index("Remark") + 1  # 1-based
    user_col = headers.index("ODL User") + 1

    # Build lookup
    remarks_map = {}
    for r in results:
        remarks_map[r["username"]] = f"[OneDrive] {r['cleanup_status']}: {r['remark'][:150]}"

    updated = 0
    for row_idx in range(2, ws.max_row + 1):
        username = ws.cell(row=row_idx, column=user_col).value
        if username and str(username).strip() in remarks_map:
            existing = ws.cell(row=row_idx, column=remark_col).value or ""
            new_remark = remarks_map[str(username).strip()]
            if existing:
                ws.cell(row=row_idx, column=remark_col).value = f"{existing} | {new_remark}"
            else:
                ws.cell(row=row_idx, column=remark_col).value = new_remark
            updated += 1

    wb.save(EXCEL_FILE)
    wb.close()
    print(f"\nUpdated {updated} remarks in RUNNING sheet of {EXCEL_FILE}")


def main():
    accounts = load_running_accounts()
    print(f"Loaded {len(accounts)} RUNNING accounts")

    if START_INDEX > 0:
        accounts = accounts[START_INDEX:]
    if MAX_COUNT > 0:
        accounts = accounts[:MAX_COUNT]

    total = len(accounts)
    print(f"Processing {total} accounts (cleanup OneDrive, keep '{KEEP_FOLDER}')")
    print(f"Results: {RESULTS_FILE}")

    # Load already-done usernames to skip
    done_users = set()
    if os.path.exists(RESULTS_FILE):
        with open(RESULTS_FILE, 'r', encoding='utf-8') as f:
            for row in csv.DictReader(f):
                done_users.add(row['username'])
        print(f'Skipping {len(done_users)} already-done accounts')

    results = []
    for i, acc in enumerate(accounts, 1):
        if acc['ODL User'] in done_users:
            print(f'[{i}/{total}] SKIP (already done): {acc["ODL User"]}')
            continue
        r = process_account(acc, i, total)
        results.append(r)
        save_result_row(r)

    # Summary
    print(f"\n{'='*60}")
    print(f"CLEANUP SUMMARY")
    print(f"{'='*60}")
    from collections import Counter
    status_counts = Counter(r["cleanup_status"] for r in results)
    for s, c in sorted(status_counts.items()):
        print(f"  {s}: {c}")
    total_extra = sum(r["extra_items"] for r in results)
    total_deleted = sum(r["deleted"] for r in results)
    print(f"  Total extra items found: {total_extra}")
    print(f"  Total items deleted: {total_deleted}")
    print(f"  TOTAL accounts: {total}")

    # Update Excel
    update_excel_remarks(results)
    print(f"\nResults saved to: {RESULTS_FILE}")


if __name__ == "__main__":
    main()
