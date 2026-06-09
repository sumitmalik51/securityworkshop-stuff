#!/usr/bin/env python3
"""
IRM (Insider Risk Management) Policy Screenshot Script.
Navigates to IRM Policies page, takes screenshots, generates markdown report.

Usage:
  python check_irm_policy.py Batch1 0 0 5    # Batch1, start=0, count=0(all), parallel=5
  python check_irm_policy.py All 0 10 3      # All sheet, first 10, 3 parallel
"""

import time
import sys
import os
import csv
import openpyxl
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from playwright.sync_api import sync_playwright

SPN_FILE = r"C:\certs\spns.xlsx"
SPN_SHEET = sys.argv[1] if len(sys.argv) > 1 else "Batch1"
START_INDEX = int(sys.argv[2]) if len(sys.argv) > 2 else 0
MAX_COUNT = int(sys.argv[3]) if len(sys.argv) > 3 else 0
PARALLEL = int(sys.argv[4]) if len(sys.argv) > 4 else 5

RESULTS_FILE = rf"C:\certs\irm_results_{SPN_SHEET}.csv"
SCREENSHOTS_DIR = rf"C:\certs\irm_screenshots_{SPN_SHEET}"
REPORT_FILE = rf"C:\certs\irm_report_{SPN_SHEET}.md"


def load_accounts():
    accounts = []
    wb = openpyxl.load_workbook(SPN_FILE, read_only=True)
    ws = wb[SPN_SHEET]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    user_col = headers.index("odluser")
    pwd_col = headers.index("odlpassword")
    tid_col = headers.index("TenantId")
    for row in ws.iter_rows(min_row=2, values_only=True):
        username = row[user_col]
        password = row[pwd_col]
        tenant_id = row[tid_col]
        if username and password and tenant_id:
            accounts.append((str(username).strip(), str(password).strip(), str(tenant_id).strip()))
    wb.close()
    return accounts


def close_popups(page):
    for _ in range(5):
        closed = False
        for sel in [
            'button[aria-label="Close"]', 'button[aria-label="Dismiss"]',
            'button:has-text("Close")', 'button:has-text("Dismiss")',
            'button:has-text("Got it")', 'button:has-text("Skip")',
            'button:has-text("OK")', 'button:has-text("No thanks")',
            'button:has-text("Maybe later")', 'button:has-text("Not now")',
            'button:has-text("Get started")',
            'div[role="dialog"] button[aria-label="Close"]',
        ]:
            try:
                btn = page.locator(sel)
                if btn.count() > 0 and btn.first.is_visible():
                    btn.first.click(timeout=2000, force=True)
                    closed = True
                    time.sleep(1)
            except Exception:
                pass
        if not closed:
            break
    try:
        page.evaluate('document.querySelectorAll(".fui-DialogSurface__backdrop").forEach(e => e.remove())')
        page.evaluate('document.querySelectorAll("[aria-hidden=true]").forEach(e => { if(e.classList.contains("fui-DialogSurface__backdrop")) e.remove() })')
    except Exception:
        pass


def check_irm(username, password, tenant_id, index, total):
    print(f"\n=== [{index}/{total}] {username} ===")
    result = {
        "username": username,
        "tenant_id": tenant_id,
        "status": "ERROR",
        "policies": "",
        "policy_count": "",
        "page_text": "",
        "screenshot": "",
        "checked_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }

    for attempt in range(3):
        if attempt > 0:
            print(f"  Retry {attempt+1}/3")
            time.sleep(5)
        browser = None
        try:
            with sync_playwright() as p:
                browser = p.chromium.launch(headless=False)
                context = browser.new_context(viewport={"width": 1920, "height": 1080})
                page = context.new_page()

                # Login
                for _nav in range(3):
                    try:
                        page.goto("https://purview.microsoft.com", timeout=120000)
                        break
                    except Exception:
                        if _nav == 2:
                            raise
                        time.sleep(10)
                time.sleep(3)
                try:
                    page.wait_for_load_state("load", timeout=30000)
                except Exception:
                    pass

                page.fill('input[type="email"]', username)
                page.click('input[type="submit"]')
                time.sleep(4)

                try:
                    pw_link = page.locator('a:has-text("Use your password instead"), a:has-text("password"), #idA_PWD_SwitchToPassword')
                    pw_link.first.click(timeout=8000)
                    time.sleep(3)
                except Exception:
                    pass

                page.fill('input[type="password"]', password)
                page.click('input[type="submit"]')
                time.sleep(5)

                try:
                    yes_btn = page.locator('input[type="submit"][value="Yes"], button:has-text("Yes")')
                    yes_btn.first.click(timeout=8000)
                    time.sleep(3)
                except Exception:
                    pass

                try:
                    page.wait_for_load_state("load", timeout=30000)
                except Exception:
                    pass
                time.sleep(3)
                close_popups(page)

                # Navigate to IRM Policies page
                irm_url = f"https://purview.microsoft.com/insiderriskmgmt/policiespage?tid={tenant_id}"
                print(f"  Navigating to IRM Policies...")
                for _nav in range(3):
                    try:
                        page.goto(irm_url, wait_until="domcontentloaded", timeout=120000)
                        break
                    except Exception:
                        if _nav == 2:
                            raise
                        time.sleep(10)
                time.sleep(10)

                # Click "Get started" first if it appears (must click before dismissing)
                try:
                    gs = page.locator('button:has-text("Get started")')
                    if gs.count() > 0 and gs.first.is_visible():
                        gs.first.click(timeout=5000, force=True)
                        print("  Clicked 'Get started'")
                        time.sleep(5)
                except Exception:
                    pass

                # Close popups multiple rounds
                for _ in range(3):
                    close_popups(page)
                    time.sleep(2)
                    try:
                        dialogs = page.locator('[role="dialog"], .ms-Dialog, .fui-DialogSurface')
                        if dialogs.count() == 0:
                            break
                    except Exception:
                        break

                # Wait for page content to fully load (up to 120s)
                # Must wait until "Loading..." text and shimmer rows disappear
                print("  Waiting for policies to load...")
                for _w in range(24):
                    close_popups(page)
                    try:
                        is_loading = False
                        # Check for "Loading..." text
                        loading_text = page.locator('text=/Loading/i')
                        if loading_text.count() > 0:
                            is_loading = True
                        # Check for shimmer/spinner placeholders
                        shimmers = page.locator('.ms-Shimmer-container, .ms-Spinner, [class*="shimmer"], [class*="Shimmer"], [class*="spinner"], [class*="loading"]')
                        if shimmers.count() > 0:
                            is_loading = True
                        if is_loading:
                            if _w % 4 == 0:
                                print(f"  Still loading... ({_w*5}s)")
                            time.sleep(5)
                            continue
                        # No loading indicators — check if data is present
                        has_error = page.locator('text=/don.*have.*permission|not authorized|access denied|something went wrong/i').count() > 0
                        if has_error:
                            break
                        # Data loaded (or empty)
                        break
                    except Exception:
                        pass
                    time.sleep(5)

                # Extra wait after loading completes for rendering
                time.sleep(3)

                close_popups(page)
                time.sleep(2)

                # Verify we're on the IRM page
                current_url = page.url
                if "insiderriskmgmt" in current_url.lower() or "policiespage" in current_url.lower():
                    print(f"  Confirmed on IRM Policies page")
                else:
                    print(f"  [WARN] URL: {current_url}")

                # Check for errors/permission issues
                try:
                    error_el = page.locator('text=/don.*have.*permission|not authorized|access denied|something went wrong/i')
                    if error_el.count() > 0:
                        err_text = error_el.first.inner_text()[:100]
                        result["status"] = "NO_ACCESS"
                        result["page_text"] = err_text
                        print(f"  NO_ACCESS: {err_text}")
                except Exception:
                    pass

                # Try to read policy info from the page
                if result["status"] != "NO_ACCESS":
                    try:
                        # Look for policy name links/cells (e.g. "Lab - IRM ...")
                        policy_links = page.locator('[role="gridcell"] a, [role="gridcell"] button, [data-automationid="DetailsRowCell"] a, [data-automationid="DetailsRowCell"] button, td a')
                        policy_names = []
                        count = policy_links.count()
                        for i in range(min(count, 20)):
                            try:
                                text = policy_links.nth(i).inner_text(timeout=2000).strip()
                                if text and len(text) > 3:
                                    policy_names.append(text)
                            except Exception:
                                pass

                        # Fallback: try getting text from "Policy name" column cells
                        if not policy_names:
                            rows = page.locator('[role="row"]')
                            row_count = rows.count()
                            for i in range(1, min(row_count, 20)):
                                try:
                                    cells = rows.nth(i).locator('[role="gridcell"]')
                                    if cells.count() >= 2:
                                        # Policy name is typically the second cell (first is checkbox)
                                        text = cells.nth(1).inner_text(timeout=2000).strip()
                                        if text and len(text) > 3:
                                            policy_names.append(text.split('\n')[0].strip())
                                except Exception:
                                    pass

                        if policy_names:
                            result["status"] = "HAS_POLICIES"
                            result["policies"] = " | ".join(policy_names[:10])
                            result["policy_count"] = str(len(policy_names))
                            print(f"  HAS_POLICIES: {len(policy_names)} found")
                        else:
                            # No policy names found - check for "no policies" text
                            no_policy = page.locator('text=/no policies|create.*policy/i')
                            if no_policy.count() > 0:
                                result["status"] = "NO_POLICIES"
                                print(f"  NO_POLICIES")
                            else:
                                result["status"] = "LOADED"
                                print(f"  Page loaded (no table found)")
                    except Exception as e:
                        print(f"  [WARN] Policy check: {e}")
                        if result["status"] == "ERROR":
                            result["status"] = "LOADED"

                # Take screenshot
                try:
                    safe_name = username.split('@')[0]
                    screenshot_path = os.path.join(SCREENSHOTS_DIR, f"{safe_name}.png")
                    page.screenshot(path=screenshot_path, full_page=True)
                    result["screenshot"] = screenshot_path
                    print(f"  Screenshot saved")
                except Exception as e:
                    print(f"  [WARN] Screenshot failed: {e}")

        except Exception as e:
            print(f"  ERROR (attempt {attempt+1}): {e}")
            result["status"] = "ERROR"
        finally:
            try:
                if browser:
                    browser.close()
            except Exception:
                pass

        if result["status"] not in ("ERROR",):
            break

    print(f"  RESULT: {result['status']}")
    return result


def load_existing_results():
    existing = {}
    if os.path.exists(RESULTS_FILE):
        try:
            with open(RESULTS_FILE, "r", newline="", encoding="utf-8") as f:
                for row in csv.DictReader(f):
                    existing[row["username"]] = row
            print(f"Loaded {len(existing)} existing results")
        except Exception:
            pass
    return existing


def save_results(all_results):
    fieldnames = ["username", "tenant_id", "status", "policy_count", "policies",
                  "page_text", "screenshot", "checked_at"]
    with open(RESULTS_FILE, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(all_results.values())


def generate_report(all_results):
    all_vals = sorted(all_results.values(), key=lambda r: r.get("status", ""))
    with open(REPORT_FILE, "w", encoding="utf-8") as f:
        f.write(f"# IRM Policy Report - {SPN_SHEET}\n\n")
        f.write(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")

        has_policies = sum(1 for r in all_vals if r["status"] == "HAS_POLICIES")
        no_policies = sum(1 for r in all_vals if r["status"] == "NO_POLICIES")
        no_access = sum(1 for r in all_vals if r["status"] == "NO_ACCESS")
        loaded = sum(1 for r in all_vals if r["status"] == "LOADED")
        errors = sum(1 for r in all_vals if r["status"] == "ERROR")

        f.write(f"| Status | Count |\n")
        f.write(f"|--------|-------|\n")
        f.write(f"| HAS_POLICIES | {has_policies} |\n")
        f.write(f"| NO_POLICIES | {no_policies} |\n")
        f.write(f"| LOADED | {loaded} |\n")
        f.write(f"| NO_ACCESS | {no_access} |\n")
        f.write(f"| ERROR | {errors} |\n")
        f.write(f"| **TOTAL** | **{len(all_vals)}** |\n\n")
        f.write(f"---\n\n")

        for r in all_vals:
            user = r.get("username", "")
            status = r.get("status", "")
            tid = r.get("tenant_id", "")
            checked = r.get("checked_at", "")
            screenshot = r.get("screenshot", "")

            f.write(f"## {user}\n\n")
            f.write(f"- **Status:** {status}\n")
            f.write(f"- **Tenant ID:** {tid}\n")
            f.write(f"- **Checked at:** {checked}\n")
            if r.get("policy_count"):
                f.write(f"- **Policy Count:** {r['policy_count']}\n")
            if r.get("policies"):
                f.write(f"- **Policies:** {r['policies']}\n")
            if r.get("page_text"):
                f.write(f"- **Page Text:** {r['page_text']}\n")

            if screenshot and os.path.exists(screenshot):
                rel_path = f"irm_screenshots_{SPN_SHEET}/{os.path.basename(screenshot)}"
                f.write(f"\n![{user}]({rel_path})\n")
            else:
                f.write(f"\n*No screenshot available*\n")
            f.write(f"\n---\n\n")

    print(f"Report saved to: {REPORT_FILE}")


def main():
    os.makedirs(SCREENSHOTS_DIR, exist_ok=True)
    accounts = load_accounts()
    print(f"Loaded {len(accounts)} accounts from '{SPN_SHEET}'")

    all_results = load_existing_results()

    if START_INDEX > 0:
        accounts = accounts[START_INDEX:]
    if MAX_COUNT > 0:
        accounts = accounts[:MAX_COUNT]

    total = len(accounts)
    print(f"Checking {total} accounts (parallel={PARALLEL})")
    if total == 0:
        print("Nothing to check!")
        return

    done_count = [0]

    def _handle(r):
        all_results[r["username"]] = r
        done_count[0] += 1
        print(f"  Progress: {done_count[0]}/{total}")
        save_results(all_results)

    if PARALLEL <= 1:
        for i, (u, p, t) in enumerate(accounts):
            _handle(check_irm(u, p, t, i + 1, total))
    else:
        with ThreadPoolExecutor(max_workers=PARALLEL) as ex:
            futs = {ex.submit(check_irm, u, p, t, i + 1, total): i
                    for i, (u, p, t) in enumerate(accounts)}
            for f in as_completed(futs):
                _handle(f.result())

    # Summary
    vals = list(all_results.values())
    has_policies = sum(1 for r in vals if r["status"] == "HAS_POLICIES")
    no_policies = sum(1 for r in vals if r["status"] == "NO_POLICIES")

    print(f"\n{'='*60}")
    print(f"IRM RESULTS — {SPN_SHEET}")
    print(f"{'='*60}")
    print(f"  HAS_POLICIES:  {has_policies}")
    print(f"  NO_POLICIES:   {no_policies}")
    print(f"  LOADED:        {sum(1 for r in vals if r['status'] == 'LOADED')}")
    print(f"  NO_ACCESS:     {sum(1 for r in vals if r['status'] == 'NO_ACCESS')}")
    print(f"  ERROR:         {sum(1 for r in vals if r['status'] == 'ERROR')}")
    print(f"  TOTAL:         {len(vals)}")
    print(f"\nResults: {RESULTS_FILE}")

    generate_report(all_results)


if __name__ == "__main__":
    main()
