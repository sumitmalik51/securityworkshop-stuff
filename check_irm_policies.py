#!/usr/bin/env python3
"""
Check Insider Risk Management (IRM) Policy warnings across all tenants.
Navigates to Purview > Insider Risk Management > Policies, captures:
  - Number of policy warnings, recommendations, healthy policies
  - Per-policy: name, status text (warnings/recommendations), users in scope, active alerts
  - Screenshot of the policies page
Outputs CSV + screenshots. Runs one account at a time for reliability.

Usage:
  python check_irm_policies.py <BatchSheet> [start_index] [max_count] [parallel]
  python check_irm_policies.py Batch1         # all accounts in Batch1
  python check_irm_policies.py Batch1 0 5     # first 5 accounts
"""

import time
import sys
import os
import csv
import re
import openpyxl
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from playwright.sync_api import sync_playwright

SPN_FILE = r"C:\certs\spns.xlsx"
SPN_SHEET = sys.argv[1] if len(sys.argv) > 1 else "Batch1"
START_INDEX = int(sys.argv[2]) if len(sys.argv) > 2 else 0
MAX_COUNT = int(sys.argv[3]) if len(sys.argv) > 3 else 0
PARALLEL = int(sys.argv[4]) if len(sys.argv) > 4 else 1

RESULTS_FILE = rf"C:\certs\irm_policies_{SPN_SHEET}.csv"
SCREENSHOTS_DIR = rf"C:\certs\screenshots_irm_{SPN_SHEET}"

# IRM Policies URL pattern
IRM_URL = "https://purview.microsoft.com/insider-risk-management/policies?tid={tid}"


def load_accounts():
    accounts = []
    wb = openpyxl.load_workbook(SPN_FILE, read_only=True)
    ws = wb[SPN_SHEET]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    user_col = headers.index("odluser")
    pwd_col = headers.index("odlpassword")
    tid_col = headers.index("TenantId")
    for row in ws.iter_rows(min_row=2, values_only=True):
        username = row[user_col]
        password = row[pwd_col]
        tenant_id = row[tid_col]
        if username and password and tenant_id:
            accounts.append((str(username).strip(), str(password).strip(), str(tenant_id).strip()))
    wb.close()
    return accounts


def dismiss_dialogs(page):
    """Dismiss popups, classic portal switch, copilot panel."""
    # Classic portal switch
    for sel in [
        'a:has-text("new portal")', 'a:has-text("Switch")', 'a:has-text("Try now")',
        'button:has-text("Switch")', 'button:has-text("Try now")',
    ]:
        try:
            btn = page.locator(sel)
            if btn.count() > 0 and btn.first.is_visible():
                btn.first.click(timeout=3000)
                time.sleep(2)
        except Exception:
            pass

    # Standard dismiss
    for sel in [
        'button[aria-label="Close"]', 'button[aria-label="Dismiss"]',
        'button:has-text("Close")', 'button:has-text("Dismiss")',
        'button:has-text("Got it")', 'button:has-text("Skip")',
        'button:has-text("OK")', 'button:has-text("No thanks")',
        'button:has-text("Maybe later")', 'button:has-text("Not now")',
        'div[role="dialog"] button[aria-label="Close"]',
    ]:
        try:
            btn = page.locator(sel)
            if btn.count() > 0 and btn.first.is_visible():
                btn.first.click(timeout=2000)
                time.sleep(1)
        except Exception:
            pass

    # Copilot panel
    try:
        cp = page.locator('button[aria-label="Close Security Copilot"], button[aria-label="Close Copilot"], button[aria-label="Close panel"]')
        if cp.count() > 0 and cp.first.is_visible():
            cp.first.click(timeout=3000, force=True)
            time.sleep(1)
    except Exception:
        pass

    # Remove overlays
    try:
        page.evaluate('document.querySelectorAll(".fui-DialogSurface__backdrop, .ms-Overlay").forEach(e => e.remove())')
    except Exception:
        pass


def handle_get_started(page):
    """Handle first-time setup for fresh accounts."""
    try:
        gs = page.locator('button:has-text("Get started")')
        if gs.count() > 0 and gs.first.is_visible():
            gs.first.click(timeout=5000)
            time.sleep(10)
            try:
                ss = page.locator('button:has-text("Start setup")')
                if ss.count() > 0 and ss.first.is_visible():
                    ss.first.click(timeout=5000)
                    time.sleep(12)
            except Exception:
                pass
            try:
                cb = page.locator('button:has-text("Close")')
                if cb.count() > 0 and cb.first.is_visible():
                    cb.first.click(timeout=5000)
                    time.sleep(3)
            except Exception:
                pass
    except Exception:
        pass


def check_irm_policies(username, password, tenant_id, index, total):
    """Check IRM policies for a single tenant."""
    safe_name = username.split("@")[0]
    print(f"\n{'='*70}")
    print(f"[{index}/{total}] {username}")
    print(f"  TenantID: {tenant_id}")
    print(f"{'='*70}")

    result = {
        "username": username,
        "tenant_id": tenant_id,
        "policy_warnings": "",
        "policy_recommendations": "",
        "healthy_policies": "",
        "policies": "",  # JSON-like: name|status|users|alerts; ...
        "irm_status": "ERROR",
        "remark": "",
        "checked_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "screenshot": "",
    }

    browser = None
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=False)
            context = browser.new_context(viewport={"width": 1920, "height": 1080})
            page = context.new_page()

            # === LOGIN ===
            print("  [1/4] Logging in...")
            page.goto("https://purview.microsoft.com", timeout=120000)
            time.sleep(3)
            try:
                page.wait_for_load_state("load", timeout=30000)
            except Exception:
                pass

            # Email
            try:
                page.fill('input[type="email"]', username, timeout=15000)
                page.click('input[type="submit"]')
                time.sleep(4)
            except Exception as e:
                print(f"  [ERROR] Email: {e}")
                result["remark"] = f"Login error: {str(e)[:100]}"
                return result

            # Password link
            try:
                pw_link = page.locator('a:has-text("Use your password instead"), #idA_PWD_SwitchToPassword')
                pw_link.first.click(timeout=8000)
                time.sleep(3)
            except Exception:
                pass

            # Password
            try:
                page.fill('input[type="password"]', password)
                page.click('input[type="submit"]')
                time.sleep(5)
            except Exception as e:
                print(f"  [ERROR] Password: {e}")
                result["remark"] = f"Password error: {str(e)[:100]}"
                return result

            # Stay signed in
            try:
                yes_btn = page.locator('input[type="submit"][value="Yes"], button:has-text("Yes")')
                yes_btn.first.click(timeout=8000)
                time.sleep(3)
            except Exception:
                pass

            try:
                page.wait_for_load_state("load", timeout=30000)
            except Exception:
                pass
            time.sleep(3)

            # Check login failure
            if "login.microsoftonline.com" in page.url or "login.live.com" in page.url:
                err = page.locator('#usernameError, #passwordError, #errorText, .alert-error')
                if err.count() > 0:
                    result["remark"] = f"Login failed: {err.first.inner_text()[:100]}"
                    result["irm_status"] = "LOGIN_FAILED"
                    return result

            print("  [2/4] Dismissing popups...")
            dismiss_dialogs(page)
            handle_get_started(page)
            dismiss_dialogs(page)

            # === NAVIGATE TO IRM POLICIES ===
            print("  [3/4] Navigating to IRM Policies...")
            irm_url = IRM_URL.format(tid=tenant_id)
            page.goto(irm_url, wait_until="domcontentloaded", timeout=120000)
            try:
                page.wait_for_load_state("load", timeout=30000)
            except Exception:
                pass
            time.sleep(10)

            dismiss_dialogs(page)
            handle_get_started(page)
            dismiss_dialogs(page)

            # If redirected to compliance portal
            if "compliance.microsoft.com" in page.url:
                print("  [ACTION] Redirected to compliance portal, re-navigating...")
                page.goto(irm_url, wait_until="domcontentloaded", timeout=120000)
                time.sleep(10)
                dismiss_dialogs(page)

            # Wait for the policies page to load
            print("  [4/4] Reading policy data...")
            policies_loaded = False
            for wait in range(24):  # up to 120s
                dismiss_dialogs(page)

                # Check for permission errors
                try:
                    perm = page.locator('text=/don\'t have permission|access denied|not authorized|Forbidden/i')
                    if perm.count() > 0:
                        result["irm_status"] = "NO_PERMISSION"
                        result["remark"] = "Insufficient permissions for IRM"
                        print("  -> No permission for IRM")
                        break
                except Exception:
                    pass

                # Check if policies page loaded (look for "Policy warnings" or "Policy name" header)
                try:
                    pw_header = page.locator('text=/Policy warnings/i, text=/Policy name/i, text=/Healthy policies/i')
                    if pw_header.count() > 0:
                        policies_loaded = True
                        print(f"  -> Policies page loaded (after {wait*5}s)")
                        break
                except Exception:
                    pass

                # Check for "Create policy" button as alternative indicator
                try:
                    create_btn = page.locator('button:has-text("Create policy"), text=/Create policy/i')
                    if create_btn.count() > 0:
                        policies_loaded = True
                        print(f"  -> Policies page loaded via Create policy button (after {wait*5}s)")
                        break
                except Exception:
                    pass

                if wait % 4 == 0 and wait > 0:
                    print(f"    Waiting for IRM page... ({wait*5}s)")
                time.sleep(5)

            if policies_loaded:
                time.sleep(3)
                dismiss_dialogs(page)

                # Extract summary counts
                try:
                    # Policy warnings count
                    pw_el = page.locator('text=/Policy warnings/i').first
                    pw_parent = pw_el.locator('..')
                    pw_text = pw_parent.inner_text(timeout=5000)
                    pw_match = re.search(r'(\d+)', pw_text)
                    if pw_match:
                        result["policy_warnings"] = pw_match.group(1)
                        print(f"  Policy warnings: {result['policy_warnings']}")
                except Exception:
                    pass

                try:
                    # Policy recommendations count
                    pr_el = page.locator('text=/Policy recommendations/i').first
                    pr_parent = pr_el.locator('..')
                    pr_text = pr_parent.inner_text(timeout=5000)
                    pr_match = re.search(r'(\d+)', pr_text)
                    if pr_match:
                        result["policy_recommendations"] = pr_match.group(1)
                        print(f"  Policy recommendations: {result['policy_recommendations']}")
                except Exception:
                    pass

                try:
                    # Healthy policies count
                    hp_el = page.locator('text=/Healthy policies/i').first
                    hp_parent = hp_el.locator('..')
                    hp_text = hp_parent.inner_text(timeout=5000)
                    hp_match = re.search(r'(\d+)', hp_text)
                    if hp_match:
                        result["healthy_policies"] = hp_match.group(1)
                        print(f"  Healthy policies: {result['healthy_policies']}")
                except Exception:
                    pass

                # Extract individual policy rows
                policies_list = []
                try:
                    # Look for policy rows in the table
                    # Each row has: policy name, status, users in scope, active alerts
                    rows = page.locator('div[role="row"], tr').all()
                    for row_el in rows:
                        try:
                            row_text = row_el.inner_text(timeout=3000)
                            # Skip header rows
                            if "Policy name" in row_text or not row_text.strip():
                                continue
                            # Look for "Lab -" pattern in policy names
                            if "Lab -" in row_text or "lab -" in row_text.lower():
                                # Parse the row text - typically: "PolicyName \t status \t users \t alerts"
                                parts = [p.strip() for p in row_text.split('\t') if p.strip()]
                                if not parts:
                                    parts = [p.strip() for p in row_text.split('\n') if p.strip()]
                                if len(parts) >= 1:
                                    policy_name = parts[0] if parts else row_text[:80]
                                    status_text = parts[1] if len(parts) > 1 else ""
                                    users = parts[2] if len(parts) > 2 else ""
                                    alerts = parts[3] if len(parts) > 3 else ""
                                    policies_list.append(f"{policy_name}|{status_text}|{users}|{alerts}")
                        except Exception:
                            continue
                except Exception as e:
                    print(f"  [WARN] Could not parse policy rows: {e}")

                # Alternative: try to get policy names by locator
                if not policies_list:
                    try:
                        policy_links = page.locator('text=/Lab - IRM/i').all()
                        for pl in policy_links:
                            try:
                                name = pl.inner_text(timeout=3000)
                                # Get parent row text for full info
                                parent = pl.locator('xpath=ancestor::div[contains(@role,"row")] | ancestor::tr')
                                if parent.count() > 0:
                                    full_text = parent.first.inner_text(timeout=3000)
                                    policies_list.append(full_text.replace('\n', ' | ').replace('\t', ' | ')[:200])
                                else:
                                    policies_list.append(name)
                            except Exception:
                                continue
                    except Exception:
                        pass

                if policies_list:
                    result["policies"] = "; ".join(policies_list)
                    print(f"  Found {len(policies_list)} policies")
                    for pol in policies_list:
                        print(f"    - {pol[:100]}")

                # Determine overall status
                warnings = int(result["policy_warnings"]) if result["policy_warnings"] else 0
                if warnings > 0:
                    result["irm_status"] = "HAS_WARNINGS"
                    result["remark"] = f"{warnings} policy warning(s)"
                elif result["policy_warnings"] == "0":
                    result["irm_status"] = "HEALTHY"
                    result["remark"] = "No warnings"
                else:
                    result["irm_status"] = "LOADED"
                    result["remark"] = "Policies page loaded but could not parse warnings count"

            elif result["irm_status"] != "NO_PERMISSION":
                # Check if it's a classic dialog issue
                try:
                    classic = page.locator('text=/classic|Switch to|compliance portal/i')
                    if classic.count() > 0:
                        result["irm_status"] = "CLASSIC_DIALOG"
                        result["remark"] = "Classic portal dialog blocking"
                    else:
                        result["irm_status"] = "PAGE_BLANK"
                        result["remark"] = "IRM page did not load after 120s"
                except Exception:
                    result["irm_status"] = "PAGE_BLANK"
                    result["remark"] = "IRM page did not load"

            # Screenshot
            screenshot_path = os.path.join(SCREENSHOTS_DIR, f"{safe_name}_IRM_{result['irm_status']}.png")
            try:
                page.screenshot(path=screenshot_path, full_page=True)
                result["screenshot"] = screenshot_path
                print(f"  Screenshot: {screenshot_path}")
            except Exception as e:
                print(f"  [WARN] Screenshot failed: {e}")

    except Exception as e:
        print(f"  [ERROR] {e}")
        result["irm_status"] = "ERROR"
        result["remark"] = str(e)[:200]
    finally:
        try:
            if browser:
                browser.close()
        except Exception:
            pass

    print(f"  RESULT: {result['irm_status']} | warnings={result['policy_warnings']} | {result['remark']}")
    return result


def save_results(results):
    fieldnames = ["username", "tenant_id", "policy_warnings", "policy_recommendations",
                  "healthy_policies", "policies", "irm_status", "remark", "checked_at", "screenshot"]
    with open(RESULTS_FILE, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(results)
    print(f"\nResults saved to: {RESULTS_FILE} ({len(results)} entries)")


def main():
    os.makedirs(SCREENSHOTS_DIR, exist_ok=True)
    accounts = load_accounts()
    print(f"Loaded {len(accounts)} accounts from sheet '{SPN_SHEET}'")

    if START_INDEX > 0:
        accounts = accounts[START_INDEX:]
    if MAX_COUNT > 0:
        accounts = accounts[:MAX_COUNT]

    total = len(accounts)
    print(f"Checking {total} accounts for IRM policy warnings")

    results = []

    if PARALLEL <= 1:
        for i, (user, pwd, tid) in enumerate(accounts, 1):
            r = check_irm_policies(user, pwd, tid, i, total)
            results.append(r)
            save_results(results)
    else:
        futures = {}
        with ThreadPoolExecutor(max_workers=PARALLEL) as executor:
            for i, (user, pwd, tid) in enumerate(accounts, 1):
                f = executor.submit(check_irm_policies, user, pwd, tid, i, total)
                futures[f] = i
            for f in as_completed(futures):
                r = f.result()
                results.append(r)
                save_results(results)

    # Summary
    print(f"\n{'='*60}")
    print(f"IRM POLICY CHECK SUMMARY - {SPN_SHEET}")
    print(f"{'='*60}")
    from collections import Counter
    status_counts = Counter(r["irm_status"] for r in results)
    for s, c in sorted(status_counts.items()):
        print(f"  {s}: {c}")
    print(f"  TOTAL: {total}")

    # Show accounts with warnings
    warned = [r for r in results if r["irm_status"] == "HAS_WARNINGS"]
    if warned:
        print(f"\n  Accounts with warnings ({len(warned)}):")
        for r in warned:
            print(f"    {r['username']}: {r['policy_warnings']} warnings, {r['policy_recommendations']} recommendations")


if __name__ == "__main__":
    main()
