#!/usr/bin/env python3
"""
DSPM Posture Agent Rerun Script for RUNNING accounts.
Reads from dspm_consolidated_status.xlsx RUNNING sheet (or spns.xlsx batch sheet).

Flow:
1. Login -> navigate to DSPM Asset Explorer -> Posture Agent tab
2. Check prompt state:
   - If COMPLETED -> skip, move on
   - If RUNNING -> Stop Generating -> Confirm popup -> Rerun -> Confirm -> Refresh
   - If EMPTY/FAILED/no data source -> Add data sources -> Start prompt -> Stop -> Rerun -> Confirm -> Refresh
   - If HAS_HISTORY -> click entry -> check state -> stop+rerun or fresh start

Usage:
  python dspm_rerun.py                           # all RUNNING accounts from consolidated Excel
  python dspm_rerun.py --start 0 --count 5       # first 5 only
  python dspm_rerun.py --user odl_user_2227202@otuwacne104140.onmicrosoft.com --pass ojro09JCR*tW --tap y@uC$kru --tenant 509053f4-22dc-4dde-b965-5d669e8a1824
"""

import time
import sys
import os
import csv
import argparse
import openpyxl
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from playwright.sync_api import sync_playwright

CONSOLIDATED_EXCEL = r"C:\certs\dspm_consolidated_status.xlsx"
SPN_EXCEL = r"C:\certs\spns.xlsx"
RESULTS_FILE = r"C:\certs\dspm_rerun_results.csv"
PROMPT_TEXT = "Get files with sensitive types password and credentials"
SCREENSHOT_DIR = r"C:\certs\screenshots"


def parse_args():
    parser = argparse.ArgumentParser(description="DSPM Posture Agent Rerun")
    parser.add_argument("--user", help="Single user to process")
    parser.add_argument("--pass", dest="password", help="Password")
    parser.add_argument("--tap", help="TAP/secondary password")
    parser.add_argument("--tenant", help="Tenant ID")
    parser.add_argument("--start", type=int, default=0, help="Start index (0-based)")
    parser.add_argument("--count", type=int, default=0, help="Max accounts to process (0=all)")
    parser.add_argument("--parallel", type=int, default=5, help="Number of parallel browser instances (default=5)")
    parser.add_argument("--source", default="consolidated", choices=["consolidated", "batch"],
                        help="Source: consolidated (RUNNING sheet) or batch (spns.xlsx)")
    parser.add_argument("--sheet", default="RUNNING", help="Sheet name in consolidated Excel (RUNNING, NOT_FOUND, etc.)")
    parser.add_argument("--batch", default="Batch1", help="Batch sheet name if source=batch")
    parser.add_argument("--retry", default="", help="Retry accounts with this status from results CSV (e.g. STOP_FAILED)")
    parser.add_argument("--no-close", dest="no_close", action="store_true", help="Keep browsers open after processing (for manual inspection)")
    return parser.parse_args()


def load_accounts_consolidated(sheet="RUNNING"):
    wb = openpyxl.load_workbook(CONSOLIDATED_EXCEL, read_only=True)
    sheets = [s.strip() for s in sheet.split(",")]
    accounts = []
    for sh in sheets:
        ws = wb[sh]
        print(f"Reading from consolidated sheet: {sh}")
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            d = {h: (str(v).strip() if v else "") for h, v in zip(headers, row)}
            accounts.append({
                "username": d["ODL User"],
                "password": d["ODL Password"],
                "tap": d.get("TAP", ""),
                "tenant_id": d["Tenant ID"],
                "batch": d.get("Batch", ""),
            })
    wb.close()
    print(f"Total accounts from {', '.join(sheets)}: {len(accounts)}")
    return accounts


def load_accounts_batch(sheet):
    wb = openpyxl.load_workbook(SPN_EXCEL, read_only=True)
    ws = wb[sheet]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    accounts = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = {h: (str(v).strip() if v else "") for h, v in zip(headers, row)}
        if not d.get("odluser"):
            continue
        accounts.append({
            "username": d["odluser"],
            "password": d["odlpassword"],
            "tap": d.get("TAP", ""),
            "tenant_id": d["TenantId"],
            "batch": sheet,
        })
    wb.close()
    return accounts


def load_accounts_retry(status_filter):
    """Load accounts whose LATEST status in results CSV matches status_filter, lookup creds from consolidated Excel.
    Use status_filter='ALL' to retry all non-success accounts."""
    SUCCESS_STATUSES = {"RERUN_SUCCESS", "SKIPPED_COMPLETED"}
    # Get latest status per username (last entry wins since we append)
    latest_status = {}
    with open(RESULTS_FILE, "r", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            latest_status[row["username"]] = row["rerun_status"]
    if status_filter.upper() == "ALL":
        failed_users = {u for u, s in latest_status.items() if s not in SUCCESS_STATUSES}
        label = "ALL failed"
    else:
        failed_users = {u for u, s in latest_status.items() if s == status_filter}
        label = status_filter
    # Lookup creds from consolidated Excel
    all_accounts = load_accounts_consolidated()
    accounts = [a for a in all_accounts if a["username"] in failed_users]
    print(f"Retry: found {len(accounts)}/{len(failed_users)} {label} accounts with creds")
    return accounts


def close_popups(page):
    """Close welcome/tour/copilot popups."""
    try:
        btn = page.locator('button:has-text("Get started")')
        if btn.count() > 0 and btn.first.is_visible():
            btn.first.click(timeout=3000)
            time.sleep(2)
    except Exception:
        pass

    for _ in range(10):
        closed = False
        for sel in [
            'button[aria-label="Close"]', 'button[aria-label="Dismiss"]',
            'button:has-text("Close")', 'button:has-text("Dismiss")',
            'button:has-text("Got it")', 'button:has-text("Skip")',
            'button:has-text("OK")', 'button:has-text("No thanks")',
            'button:has-text("Maybe later")', 'button:has-text("Not now")',
            'div[role="dialog"] button[aria-label="Close"]',
        ]:
            try:
                btn = page.locator(sel)
                if btn.count() > 0 and btn.first.is_visible():
                    btn.first.click(timeout=2000)
                    closed = True
                    time.sleep(1)
            except Exception:
                pass
        if not closed:
            break

    try:
        cp = page.locator('button[aria-label="Close Security Copilot"], button[aria-label="Close Copilot"], button[aria-label="Close panel"]')
        if cp.count() > 0 and cp.first.is_visible():
            cp.first.click(timeout=3000, force=True)
            time.sleep(1)
    except Exception:
        pass

    try:
        page.evaluate('document.querySelectorAll(".fui-DialogSurface__backdrop").forEach(e => e.remove())')
    except Exception:
        pass


def login(page, username, password):
    """Login to Purview portal."""
    page.goto("https://purview.microsoft.com", wait_until="domcontentloaded", timeout=60000)
    time.sleep(5)

    # Wait for email input to be ready
    page.wait_for_selector('input[type="email"]', timeout=30000)
    page.fill('input[type="email"]', username)
    page.click('input[type="submit"]')
    time.sleep(4)

    # Password option if needed
    try:
        pw_link = page.locator('a:has-text("Use your password instead"), a:has-text("password"), #idA_PWD_SwitchToPassword')
        pw_link.first.click(timeout=8000)
        time.sleep(3)
    except Exception:
        pass

    try:
        page.wait_for_selector('input[type="password"]', timeout=10000)
        page.fill('input[type="password"]', password)
        page.click('input[type="submit"]')
    except Exception as e:
        print(f"  [WARN] Password issue: {e}")
    time.sleep(4)

    # Stay signed in - Yes
    try:
        yes_btn = page.locator('input[type="submit"][value="Yes"], button:has-text("Yes")')
        yes_btn.first.click(timeout=15000)
    except Exception:
        pass
    time.sleep(5)

    try:
        page.wait_for_load_state("load", timeout=30000)
    except Exception:
        pass
    time.sleep(5)
    close_popups(page)
    time.sleep(2)


def navigate_to_posture_agent(page, tenant_id):
    """Navigate to DSPM Asset Explorer -> Posture Agent tab."""
    dspm_url = f"https://purview.microsoft.com/dspm/assetexplorer?tid={tenant_id}"
    print(f"  Navigating to DSPM Asset Explorer...")
    page.goto(dspm_url, wait_until="domcontentloaded", timeout=60000)
    time.sleep(10)
    close_popups(page)
    time.sleep(2)

    # Click "Posture agent (preview)" tab
    print("  Clicking 'Posture agent (preview)'...")
    try:
        posture = page.locator('button:has-text("Posture Agent"), button:has-text("Posture agent")')
        if posture.count() == 0:
            posture = page.get_by_text("Posture agent", exact=False)
        posture.first.click(timeout=10000, force=True)
        print("  Clicked Posture agent tab")
        time.sleep(8)
    except Exception as e:
        print(f"  [WARN] Could not click Posture agent: {e}")

    close_popups(page)
    time.sleep(2)


def handle_setup_dialogs(page):
    """Handle 'Get started' and 'Start setup' dialogs for fresh accounts."""
    clicked_get_started = False
    try:
        get_started = page.locator('button:has-text("Get started")')
        if get_started.count() > 0 and get_started.first.is_visible():
            get_started.first.click(timeout=5000)
            print("  Clicked Get started")
            clicked_get_started = True
            time.sleep(10)
    except Exception:
        pass

    try:
        start_setup = page.locator('button:has-text("Start setup")')
        if clicked_get_started:
            start_setup.first.wait_for(state="visible", timeout=15000)
        if start_setup.count() > 0 and start_setup.first.is_visible():
            start_setup.first.click(timeout=5000)
            print("  Clicked Start setup")
            time.sleep(12)
    except Exception:
        pass

    clicked_close = False
    try:
        close_btn = page.locator('button:has-text("Close")')
        if close_btn.count() > 0 and close_btn.first.is_visible():
            close_btn.first.click(timeout=5000)
            print("  Closed setup dialog")
            clicked_close = True
            time.sleep(3)
    except Exception:
        pass

    if clicked_close or clicked_get_started:
        try:
            page.reload(wait_until="domcontentloaded", timeout=30000)
            print("  Refreshed after setup")
            time.sleep(10)
        except Exception:
            pass
        return True
    return False


def add_data_sources(page, username):
    """Add data sources: search user -> uncheck Mailboxes -> keep Sites -> Save."""
    print("  === Adding data sources ===")

    # Handle any setup dialogs first
    did_setup = handle_setup_dialogs(page)
    close_popups(page)
    time.sleep(2)

    if did_setup:
        # Re-click Posture Agent after setup
        print("  Re-clicking Posture agent after setup...")
        try:
            posture = page.locator('button:has-text("Posture Agent"), button:has-text("Posture agent")')
            if posture.count() == 0:
                posture = page.get_by_text("Posture agent", exact=False)
            posture.first.click(timeout=10000, force=True)
            time.sleep(8)
        except Exception:
            pass
        close_popups(page)
        time.sleep(2)

    # Click "Add data sources"
    print("  Clicking 'Add data sources'...")
    try:
        page.evaluate('document.querySelectorAll(".fui-DialogSurface__backdrop").forEach(e => e.remove())')
        time.sleep(1)
        add_data = page.locator('button:has-text("Add data sources"), a:has-text("Add data sources"), [role="button"]:has-text("Add data sources")')
        if add_data.count() == 0:
            add_data = page.locator('button:has-text("Add data"), a:has-text("Add data"), [role="button"]:has-text("Add data")')
        if add_data.count() == 0:
            add_data = page.get_by_text("Add data sources", exact=False)
        add_data.first.click(timeout=10000, force=True)
        print("  Clicked Add data sources")
        time.sleep(5)
    except Exception as e:
        print(f"  [WARN] Could not click Add data sources: {e}")
        # JS fallback
        try:
            page.evaluate('''() => {
                const walker = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT);
                while (walker.nextNode()) {
                    if (walker.currentNode.textContent.trim().toLowerCase().includes('add data source')) {
                        walker.currentNode.parentElement.click();
                        return true;
                    }
                }
                return false;
            }''')
            print("  Clicked Add data sources via JS")
            time.sleep(5)
        except Exception:
            print("  [WARN] JS fallback also failed")
            return False

    # Handle late setup dialogs
    did_setup2 = handle_setup_dialogs(page)
    if did_setup2:
        close_popups(page)
        time.sleep(2)
        print("  Re-clicking Posture agent after late setup...")
        try:
            posture = page.locator('button:has-text("Posture Agent"), button:has-text("Posture agent")')
            if posture.count() == 0:
                posture = page.get_by_text("Posture agent", exact=False)
            posture.first.click(timeout=10000, force=True)
            time.sleep(8)
        except Exception:
            pass
        close_popups(page)
        time.sleep(2)
        print("  Re-clicking Add data sources...")
        try:
            add_data = page.locator('button:has-text("Add data sources"), a:has-text("Add data sources")')
            if add_data.count() == 0:
                add_data = page.get_by_text("Add data sources", exact=False)
            add_data.first.click(timeout=10000)
            time.sleep(5)
        except Exception:
            pass

    close_popups(page)
    time.sleep(2)

    # Search for user
    user_search = username.split("@")[0]
    print(f"  Searching for '{user_search}'...")
    try:
        search_input = page.locator('input[placeholder="Search for people, groups, or sites"]')
        search_input.fill(user_search)
        time.sleep(1)
        search_input.press("Enter")
        time.sleep(8)
    except Exception as e:
        print(f"  [WARN] Search issue: {e}")
        return False

    # Select user from results
    print("  Selecting user...")
    try:
        user_btn = page.locator(f'button:has-text("{user_search}")').last
        user_btn.wait_for(state="visible", timeout=15000)
        user_btn.click(timeout=5000)
        print(f"  Selected {user_search}")
        time.sleep(5)
    except Exception as e:
        print(f"  [WARN] Button click failed ({e}), trying alternatives...")
        try:
            email_el = page.locator(f'text="{username}"').last
            if email_el.count() > 0:
                email_el.click(timeout=5000)
                time.sleep(5)
            else:
                page.evaluate(f'''() => {{
                    const els = document.querySelectorAll("button, div[role='option'], div[role='listitem'], li");
                    for (const el of els) {{
                        if (el.textContent.includes("{user_search}")) {{
                            el.click();
                            return true;
                        }}
                    }}
                    return false;
                }}''')
                time.sleep(5)
        except Exception as e2:
            print(f"  [WARN] Could not select user: {e2}")
            return False

    handle_setup_dialogs(page)
    close_popups(page)

    # Uncheck Mailboxes, keep Sites
    print("  Unchecking Mailboxes, keeping Sites...")
    try:
        page.locator('input[type="checkbox"], [role="checkbox"]').first.wait_for(timeout=15000)
        time.sleep(2)

        checkboxes = page.locator('input[type="checkbox"], [role="checkbox"]')
        cb_count = checkboxes.count()
        print(f"  Found {cb_count} checkboxes")

        if cb_count >= 2:
            mail_cb = checkboxes.nth(0)
            if mail_cb.is_visible():
                cb_id = mail_cb.get_attribute('id')
                if cb_id:
                    label = page.locator(f'label[for="{cb_id}"]')
                    if label.count() > 0:
                        label.first.click(force=True)
                    else:
                        mail_cb.click(force=True)
                else:
                    mail_cb.click(force=True)
                print("  Unchecked Mailboxes")
            print("  Sites left checked")
        time.sleep(2)
    except Exception as e:
        print(f"  [WARN] Checkbox issue: {e}")

    # Click Save
    print("  Clicking Save...")
    try:
        save_btn = page.locator('button:has-text("Save")')
        save_btn.first.click(timeout=5000, force=True)
        print("  Clicked Save")
        time.sleep(8)
    except Exception as e:
        print(f"  [WARN] Could not click Save: {e}")
        return False

    time.sleep(10)
    print("  Data sources added successfully")
    return True


def detect_prompt_state(page):
    """
    Detect current state of the posture agent prompt.
    Returns: 'COMPLETED', 'RUNNING', 'FAILED', 'NO_DATASOURCE', 'HAS_HISTORY', 'EMPTY'
    Also returns run_count if visible.
    
    "Pick up where you left off" is ALWAYS visible - not a reliable state indicator.
    Must check for actual prompt entries or "Nothing to show yet" inside it.
    """
    state = "EMPTY"
    run_count = 0

    # Check for "Stop generating" button -> means RUNNING
    try:
        stop_btn = page.locator('button:has-text("Stop generating"), button:has-text("Stop Generating")')
        if stop_btn.count() > 0 and stop_btn.first.is_visible():
            state = "RUNNING"
            print("  Detected: RUNNING (Stop generating button visible)")
    except Exception:
        pass

    # Check for "Job estimation in progress" text -> RUNNING
    if state != "RUNNING":
        try:
            job_est = page.locator('text="Job estimation in progress"')
            if job_est.count() > 0 and job_est.first.is_visible():
                state = "RUNNING"
                print("  Detected: RUNNING (Job estimation in progress)")
        except Exception:
            pass

    # Check for completed indicators
    if state != "RUNNING":
        try:
            completed_indicators = [
                page.locator('text="The Discovery agent found"'),
                page.locator('text="View insights"'),
                page.locator('button:has-text("View insights")'),
            ]
            for ind in completed_indicators:
                if ind.count() > 0 and ind.first.is_visible():
                    state = "COMPLETED"
                    print("  Detected: COMPLETED")
                    break
        except Exception:
            pass

    # Check for failed
    if state == "EMPTY":
        try:
            failed = page.locator('text="Failed"')
            if failed.count() > 0 and failed.first.is_visible():
                state = "FAILED"
                print("  Detected: FAILED")
        except Exception:
            pass

    # Check if "Add data sources" warning/button is visible -> no data source configured
    # This takes priority - even if "Pick up where you left off" is there (it always is)
    if state == "EMPTY":
        try:
            add_ds = page.locator('button:has-text("Add data sources"), a:has-text("Add data sources")')
            if add_ds.count() > 0 and add_ds.first.is_visible():
                state = "NO_DATASOURCE"
                print("  Detected: NO_DATASOURCE (Add data sources button visible)")
        except Exception:
            pass
        # Also check for the warning text
        if state == "EMPTY":
            try:
                warning = page.locator('text="Add data sources to tell the agent"')
                if warning.count() > 0 and warning.first.is_visible():
                    state = "NO_DATASOURCE"
                    print("  Detected: NO_DATASOURCE (warning text visible)")
            except Exception:
                pass

    # Check for actual prompt entries inside "Pick up where you left off"
    # "Nothing to show yet" means no previous prompts -> EMPTY
    # An actual entry like "Sensitive Password And Credentials Files" means HAS_HISTORY
    if state == "EMPTY":
        try:
            nothing_yet = page.locator('text="Nothing to show yet"')
            if nothing_yet.count() > 0 and nothing_yet.first.is_visible():
                state = "EMPTY"
                print("  Detected: EMPTY (Nothing to show yet)")
        except Exception:
            pass

    if state == "EMPTY":
        try:
            # Look for actual prompt entry text (e.g. "Sensitive Password And Credentials Files")
            prompt_entry = page.locator('text="Sensitive Password And Credentials Files"')
            if prompt_entry.count() > 0 and prompt_entry.first.is_visible():
                state = "HAS_HISTORY"
                print("  Detected: HAS_HISTORY (found prompt entry)")
        except Exception:
            pass

    # Get run count (e.g. "4 runs" badge)
    try:
        runs_badge = page.locator('text=/\\d+ runs?/')
        if runs_badge.count() > 0 and runs_badge.first.is_visible():
            runs_text = runs_badge.first.text_content()
            import re
            m = re.search(r'(\d+)\s*runs?', runs_text)
            if m:
                run_count = int(m.group(1))
                print(f"  Run count: {run_count}")
    except Exception:
        pass

    return state, run_count


def click_stop_generating(page):
    """Click Stop Generating and confirm popup."""
    print("  Clicking 'Stop generating'...")
    try:
        stop_btn = page.locator('button:has-text("Stop generating"), button:has-text("Stop Generating")')
        stop_btn.first.click(timeout=30000, force=True)
        print("  Clicked Stop generating")
        time.sleep(3)
    except Exception as e:
        print(f"  [WARN] Could not click Stop generating: {e}")
        return False

    # Confirm popup if any
    confirm_stop_popup(page)
    time.sleep(3)
    return True


def confirm_stop_popup(page):
    """Confirm any popup that appears after Stop Generating."""
    for sel in [
        'button:has-text("Yes")',
        'button:has-text("Confirm")',
        'button:has-text("Stop")',
        'button:has-text("OK")',
        'div[role="dialog"] button:has-text("Yes")',
        'div[role="dialog"] button:has-text("Confirm")',
        'div[role="dialog"] button:has-text("Stop")',
    ]:
        try:
            btn = page.locator(sel)
            if btn.count() > 0 and btn.first.is_visible():
                btn.first.click(timeout=5000, force=True)
                print(f"  Confirmed stop popup")
                time.sleep(2)
                return True
        except Exception:
            pass
    return False


def click_rerun(page):
    """Click the Rerun button (on the right side of the prompt card)."""
    print("  Clicking 'Rerun'...")
    time.sleep(3)

    # Try multiple selectors for Rerun button
    rerun_selectors = [
        'button:has-text("Rerun")',
        'button[aria-label="Rerun"]',
        'button[title="Rerun"]',
        '[role="button"]:has-text("Rerun")',
        'button:has-text("Re-run")',
        'button:has-text("Run again")',
    ]
    for sel in rerun_selectors:
        try:
            btn = page.locator(sel)
            if btn.count() > 0 and btn.first.is_visible():
                btn.first.click(timeout=10000, force=True)
                print(f"  Clicked Rerun")
                time.sleep(5)
                return True
        except Exception:
            pass

    # Fallback: try the refresh/rerun icon button (circular arrow icon)
    try:
        # Look for icon buttons near the prompt card
        icon_btns = page.locator('button[aria-label*="run"], button[aria-label*="Run"], button[aria-label*="refresh"], button[aria-label*="Refresh"]')
        if icon_btns.count() > 0:
            icon_btns.first.click(timeout=5000, force=True)
            print("  Clicked rerun icon button")
            time.sleep(5)
            return True
    except Exception:
        pass

    print("  [WARN] Could not find Rerun button")
    return False


def click_run_and_confirm(page):
    """After Rerun populates the prompt, click Run and then Confirm."""
    print("  Waiting for prompt to be populated...")
    time.sleep(3)

    # The prompt should already be filled from Rerun. Click Send/Run.
    send_selectors = [
        'button[aria-label="Send"]',
        'button:has-text("Send")',
        'button:has-text("Run")',
        'button[aria-label="Run"]',
    ]
    for sel in send_selectors:
        try:
            btn = page.locator(sel)
            if btn.count() > 0 and btn.first.is_visible():
                btn.first.click(timeout=10000, force=True)
                print(f"  Clicked Send/Run")
                time.sleep(5)
                break
        except Exception:
            pass

    # Wait for Confirm button and click it
    print("  Waiting for Confirm...")
    try:
        confirm_btn = page.locator('button:has-text("Confirm")')
        confirm_btn.wait_for(state="visible", timeout=90000)
        time.sleep(2)
        confirm_btn.click(timeout=5000, force=True)
        print("  Clicked Confirm")
        time.sleep(5)
        return True
    except Exception as e:
        print(f"  [WARN] Confirm issue: {e}")
        return False


def start_prompt_fresh(page):
    """Start a new prompt when nothing is currently running."""
    print("  Starting fresh prompt...")

    # Click "Start new prompt" if present
    try:
        start_new = page.locator('button:has-text("Start new prompt")')
        if start_new.count() > 0 and start_new.first.is_visible():
            start_new.first.click(timeout=5000)
            print("  Clicked Start new prompt")
            time.sleep(5)
    except Exception:
        pass

    # Click suggested prompt to activate prompt area
    suggested_selectors = [
        'button:has-text("Collect files")',
        'button:has-text("Show chats")',
        'button:has-text("Retrieve financial")',
        'button:has-text("Suggested prompt")',
    ]
    clicked_suggested = False
    for sel in suggested_selectors:
        try:
            btn = page.locator(sel)
            if btn.count() > 0 and btn.first.is_visible():
                btn.first.click(timeout=5000)
                print(f"  Clicked suggested prompt to activate area")
                clicked_suggested = True
                time.sleep(3)
                break
        except Exception:
            continue

    if clicked_suggested:
        page.keyboard.press("Control+a")
        time.sleep(0.5)
        page.keyboard.press("Delete")
        time.sleep(0.5)
        page.keyboard.type(PROMPT_TEXT, delay=50)
        print(f"  Typed prompt: {PROMPT_TEXT}")
        time.sleep(3)
    else:
        try:
            prompt_box = page.get_by_placeholder("Describe what you're looking for")
            prompt_box.click(timeout=10000)
            time.sleep(1)
            page.keyboard.type(PROMPT_TEXT, delay=50)
            print(f"  Typed prompt directly")
            time.sleep(3)
        except Exception as e:
            print(f"  [WARN] Could not type prompt: {e}")
            return False

    # Click Send
    send_btn = page.locator('button[aria-label="Send"], button:has-text("Send")')
    if send_btn.count() > 0:
        time.sleep(2)
        send_btn.first.click(force=True, timeout=5000)
        print("  Clicked Send")
    else:
        page.keyboard.press("Enter")
        print("  Pressed Enter to send")

    # Wait for Confirm
    print("  Waiting for Confirm...")
    try:
        confirm_btn = page.locator('button:has-text("Confirm")')
        confirm_btn.wait_for(state="visible", timeout=90000)
        time.sleep(2)
        confirm_btn.click(timeout=5000, force=True)
        print("  Clicked Confirm")
    except Exception as e:
        print(f"  [WARN] Confirm issue: {e}")
        return False

    time.sleep(5)
    return True


def take_screenshot(page, username, suffix=""):
    """Take a screenshot of the current page."""
    os.makedirs(SCREENSHOT_DIR, exist_ok=True)
    user_short = username.split("@")[0]
    fname = f"{user_short}_rerun{suffix}.png"
    fpath = os.path.join(SCREENSHOT_DIR, fname)
    try:
        page.screenshot(path=fpath, full_page=False)
        print(f"  Screenshot: {fname}")
    except Exception as e:
        print(f"  [WARN] Screenshot failed: {e}")
    return fname


def process_account(acc, index, total, no_close=False):
    username = acc["username"]
    password = acc["password"]
    tenant_id = acc["tenant_id"]
    batch = acc.get("batch", "")

    print(f"\n{'='*60}")
    print(f"[{index}/{total}] {username} ({batch})")
    print(f"{'='*60}")

    result = {
        "batch": batch,
        "username": username,
        "tenant_id": tenant_id,
        "action": "",
        "prompt_state": "",
        "run_count": 0,
        "rerun_status": "ERROR",
        "remark": "",
        "screenshot": "",
        "checked_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }

    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=False)
            context = browser.new_context()
            page = context.new_page()

            # Login
            login(page, username, password)

            # Navigate to Posture Agent tab
            navigate_to_posture_agent(page, tenant_id)

            # Detect current state
            state, run_count = detect_prompt_state(page)
            result["prompt_state"] = state
            result["run_count"] = run_count

            # Take screenshot of initial state
            take_screenshot(page, username, "_before")

            if state == "COMPLETED":
                print("  SKIP: Prompt already COMPLETED")
                result["action"] = "SKIP"
                result["rerun_status"] = "SKIPPED_COMPLETED"
                result["remark"] = "Prompt already completed, skipped"
                take_screenshot(page, username, "_completed")

            elif state == "RUNNING":
                print(f"  RUNNING detected (runs={run_count}). Will Stop + Rerun.")
                result["action"] = "STOP_AND_RERUN"

                # Step 1: Stop generating
                stopped = click_stop_generating(page)
                time.sleep(5)

                if stopped:
                    # Step 2: Click Rerun
                    reran = click_rerun(page)
                    if reran:
                        # Step 3: Click Run/Send + Confirm
                        confirmed = click_run_and_confirm(page)
                        if confirmed:
                            print("  Waiting 60s for job to start...")
                            time.sleep(60)
                            # Refresh page
                            try:
                                page.reload(wait_until="domcontentloaded", timeout=30000)
                                print("  Refreshed page")
                                time.sleep(10)
                            except Exception:
                                pass
                            result["rerun_status"] = "RERUN_SUCCESS"
                            result["remark"] = f"Stopped running job (runs={run_count}), rerun submitted"
                        else:
                            result["rerun_status"] = "CONFIRM_FAILED"
                            result["remark"] = "Stopped, rerun clicked but confirm failed"
                    else:
                        result["rerun_status"] = "RERUN_BUTTON_FAILED"
                        result["remark"] = "Stopped generating but could not find Rerun button"
                else:
                    result["rerun_status"] = "STOP_FAILED"
                    result["remark"] = "Could not click Stop Generating"

                take_screenshot(page, username, "_after")

            elif state in ("EMPTY", "HAS_HISTORY", "FAILED", "NO_DATASOURCE"):
                print(f"  State={state}. Will handle accordingly.")
                result["action"] = "FRESH_STOP_RERUN"

                # If HAS_HISTORY, click on the existing prompt entry first
                if state == "HAS_HISTORY":
                    print("  Clicking existing prompt entry...")
                    try:
                        prompt_entry = page.locator('text="Sensitive Password And Credentials Files"')
                        if prompt_entry.count() > 0:
                            prompt_entry.first.click(timeout=5000)
                            print("  Clicked existing prompt entry")
                            time.sleep(5)

                            state2, run_count2 = detect_prompt_state(page)
                            result["prompt_state"] = f"{state}->{state2}"
                            result["run_count"] = run_count2

                            if state2 == "COMPLETED":
                                print("  SKIP: After clicking entry, it's COMPLETED")
                                result["action"] = "SKIP"
                                result["rerun_status"] = "SKIPPED_COMPLETED"
                                result["remark"] = "Clicked history entry -> already completed"
                                take_screenshot(page, username, "_completed")
                                if no_close:
                                    input(f"  [{username.split('@')[0]}] Browser open. Press Enter to close...")
                                browser.close()
                                return result

                            if state2 == "RUNNING":
                                print(f"  Already RUNNING (runs={run_count2}). Will Stop + Rerun.")
                                result["action"] = "STOP_AND_RERUN"
                                stopped = click_stop_generating(page)
                                time.sleep(5)
                                if stopped:
                                    reran = click_rerun(page)
                                    if reran:
                                        confirmed = click_run_and_confirm(page)
                                        if confirmed:
                                            time.sleep(60)
                                            try:
                                                page.reload(wait_until="domcontentloaded", timeout=30000)
                                                time.sleep(10)
                                            except Exception:
                                                pass
                                            result["rerun_status"] = "RERUN_SUCCESS"
                                            result["remark"] = f"History entry was running (runs={run_count2}), stopped + rerun"
                                        else:
                                            result["rerun_status"] = "CONFIRM_FAILED"
                                            result["remark"] = "Stopped history entry but confirm failed"
                                    else:
                                        result["rerun_status"] = "RERUN_BUTTON_FAILED"
                                        result["remark"] = "Stopped but no Rerun button"
                                else:
                                    result["rerun_status"] = "STOP_FAILED"
                                    result["remark"] = "Could not stop running history entry"
                                take_screenshot(page, username, "_after")
                                if no_close:
                                    input(f"  [{username.split('@')[0]}] Browser open. Press Enter to close...")
                                browser.close()
                                return result
                    except Exception as e:
                        print(f"  [WARN] Could not click history entry: {e}")

                # For EMPTY, FAILED, or NO_DATASOURCE: need to add data sources first
                if state in ("EMPTY", "FAILED", "NO_DATASOURCE"):
                    print("  No data source / prompt found. Adding data sources first...")
                    result["action"] = "ADD_DATASOURCE_STOP_RERUN"
                    ds_ok = add_data_sources(page, username)
                    if not ds_ok:
                        result["rerun_status"] = "DATASOURCE_FAILED"
                        result["remark"] = "Could not add data sources"
                        take_screenshot(page, username, "_ds_failed")
                        if no_close:
                            input(f"  [{username.split('@')[0]}] Browser open. Press Enter to close...")
                        browser.close()
                        return result

                # Start a fresh prompt
                started = start_prompt_fresh(page)
                if started:
                    # Wait for it to start running
                    print("  Waiting 30s for prompt to start running...")
                    time.sleep(30)

                    # Now stop it
                    stopped = click_stop_generating(page)
                    time.sleep(5)

                    if stopped:
                        # Click Rerun
                        reran = click_rerun(page)
                        if reran:
                            confirmed = click_run_and_confirm(page)
                            if confirmed:
                                time.sleep(60)
                                try:
                                    page.reload(wait_until="domcontentloaded", timeout=30000)
                                    time.sleep(10)
                                except Exception:
                                    pass
                                result["rerun_status"] = "RERUN_SUCCESS"
                                result["remark"] = "Fresh prompt -> stopped -> rerun submitted"
                            else:
                                result["rerun_status"] = "CONFIRM_FAILED"
                                result["remark"] = "Fresh prompt -> stopped -> rerun but confirm failed"
                        else:
                            result["rerun_status"] = "RERUN_BUTTON_FAILED"
                            result["remark"] = "Fresh prompt -> stopped but no Rerun button"
                    else:
                        result["rerun_status"] = "STOP_FAILED"
                        result["remark"] = "Fresh prompt started but could not stop it"
                else:
                    result["rerun_status"] = "START_FAILED"
                    result["remark"] = "Could not start fresh prompt"

                take_screenshot(page, username, "_after")

            take_screenshot(page, username, "_final")
            print(f"  Result: {result['rerun_status']}")
            time.sleep(2)
            if no_close:
                input(f"  [{username.split('@')[0]}] Browser open. Press Enter to close...")
            browser.close()

    except Exception as e:
        result["rerun_status"] = "ERROR"
        result["remark"] = str(e)[:200]
        print(f"  ERROR: {e}")

    return result


def save_result_row(result):
    fieldnames = ["batch", "username", "tenant_id", "action", "prompt_state",
                  "run_count", "rerun_status", "remark", "screenshot", "checked_at"]
    file_exists = os.path.exists(RESULTS_FILE) and os.path.getsize(RESULTS_FILE) > 0
    with open(RESULTS_FILE, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        if not file_exists:
            writer.writeheader()
        writer.writerow(result)


def main():
    args = parse_args()

    # Single user mode
    if args.user:
        accounts = [{
            "username": args.user,
            "password": args.password,
            "tap": args.tap or "",
            "tenant_id": args.tenant,
            "batch": "single",
        }]
    elif args.retry:
        accounts = load_accounts_retry(args.retry)
    elif args.source == "batch":
        accounts = load_accounts_batch(args.batch)
    else:
        accounts = load_accounts_consolidated(args.sheet)

    print(f"Loaded {len(accounts)} accounts")

    if args.start > 0:
        accounts = accounts[args.start:]
    if args.count > 0:
        accounts = accounts[:args.count]

    total = len(accounts)
    print(f"Processing {total} accounts")
    print(f"Results: {RESULTS_FILE}")

    parallel = args.parallel if not args.user else 1
    print(f"Parallel: {parallel}")

    results = []
    no_close = args.no_close
    if no_close:
        print("Browsers will stay open after processing")

    if parallel <= 1:
        for i, acc in enumerate(accounts, 1):
            r = process_account(acc, i, total, no_close=no_close)
            results.append(r)
            save_result_row(r)
    else:
        def run_task(task_args):
            idx, acc = task_args
            return process_account(acc, idx, total, no_close=no_close)

        tasks = [(i, acc) for i, acc in enumerate(accounts, 1)]
        with ThreadPoolExecutor(max_workers=parallel) as executor:
            futures = {executor.submit(run_task, t): t for t in tasks}
            for future in as_completed(futures):
                try:
                    r = future.result()
                    results.append(r)
                    save_result_row(r)
                    done = len(results)
                    print(f"  Progress: {done}/{total} (last: {r['rerun_status']})")
                except Exception as e:
                    print(f"  Thread error: {e}")

    # Summary
    print(f"\n{'='*60}")
    print(f"DSPM RERUN SUMMARY")
    print(f"{'='*60}")
    from collections import Counter
    status_counts = Counter(r["rerun_status"] for r in results)
    for s, c in sorted(status_counts.items()):
        print(f"  {s}: {c}")
    print(f"  TOTAL: {total}")


if __name__ == "__main__":
    main()
