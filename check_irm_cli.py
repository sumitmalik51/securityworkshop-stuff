#!/usr/bin/env python3
"""
Check IRM Policy warnings via Security & Compliance PowerShell (CLI).
Uses Connect-IPPSSession + Get-InsiderRiskPolicy per tenant.
Much faster and more reliable than browser automation.

Usage:
  python check_irm_cli.py <BatchSheet> [start_index] [max_count]
  python check_irm_cli.py Batch1          # all accounts
  python check_irm_cli.py Batch1 0 5      # first 5 accounts
"""

import subprocess
import sys
import os
import csv
import json
import openpyxl
from datetime import datetime

SPN_FILE = r"C:\certs\spns.xlsx"
SPN_SHEET = sys.argv[1] if len(sys.argv) > 1 else "Batch1"
START_INDEX = int(sys.argv[2]) if len(sys.argv) > 2 else 0
MAX_COUNT = int(sys.argv[3]) if len(sys.argv) > 3 else 0

RESULTS_FILE = rf"C:\certs\irm_cli_{SPN_SHEET}.csv"

# PowerShell script template to run per account
PS_SCRIPT_TEMPLATE = r'''
$ErrorActionPreference = "Stop"
$WarningPreference = "SilentlyContinue"

$username = "{username}"
$password = "{password}"

Import-Module ExchangeOnlineManagement -MinimumVersion 3.0 -ErrorAction Stop

$secPwd = ConvertTo-SecureString $password -AsPlainText -Force
$cred = New-Object PSCredential($username, $secPwd)

$result = @{{
    status = "ERROR"
    policies = @()
    remark = ""
    total_policies = 0
    warning_policies = 0
    healthy_policies = 0
    recommendation_policies = 0
}}

try {{
    Connect-IPPSSession -Credential $cred -ErrorAction Stop -WarningAction SilentlyContinue 2>$null
    
    $policies = Get-InsiderRiskPolicy -ErrorAction Stop
    $result.total_policies = $policies.Count
    
    $policyList = @()
    $warningCount = 0
    $healthyCount = 0
    $recommendCount = 0
    
    foreach ($p in $policies) {{
        # Skip tenant settings policy
        if ($p.InsiderRiskScenario -eq "TenantSetting") {{ continue }}
        
        # Parse PolicyHealth (authoritative field matching UI)
        $healthStatus = ""
        $unhealthyCount = 0
        $recommendReviewCount = 0
        $validationDetails = @()
        try {{
            $ph = $p.PolicyHealth | ConvertFrom-Json
            $healthStatus = $ph.HealthStatus
            $unhealthyCount = $ph.UnhealthyCount
            $recommendReviewCount = $ph.RecommendReviewCount
            if ($ph.ValidationDetails) {{
                $validationDetails = @($ph.ValidationDetails)
            }}
        }} catch {{}}
        
        # Parse indicators for additional info
        $indicatorsEnabled = 0
        $indicatorsTotal = 0
        try {{
            $indicators = $p.Indicators | ConvertFrom-Json
            $indicatorsTotal = $indicators.Count
            $enabledList = @($indicators | Where-Object {{ $_.Enabled -eq $true }})
            $indicatorsEnabled = $enabledList.Count
        }} catch {{}}
        
        $hasWarning = ($healthStatus -eq "Unhealthy") -or ($unhealthyCount -gt 0)
        $hasRecommendation = ($healthStatus -eq "RecommendReview") -or ($recommendReviewCount -gt 0)
        if ($hasWarning) {{ $warningCount++ }}
        elseif ($hasRecommendation) {{ $recommendCount++ }}
        else {{ $healthyCount++ }}
        
        $policyInfo = @{{
            name = $p.Name
            scenario = $p.InsiderRiskScenario
            mode = $p.Mode
            health_status = $healthStatus
            unhealthy_count = $unhealthyCount
            recommend_count = $recommendReviewCount
            validation_details = ($validationDetails | ForEach-Object {{ $_ | ConvertTo-Json -Compress }}) -join "; "
            indicators_enabled = $indicatorsEnabled
            indicators_total = $indicatorsTotal
            has_warning = $hasWarning
            has_recommendation = $hasRecommendation
        }}
        $policyList += $policyInfo
    }}
    
    $result.policies = $policyList
    $result.warning_policies = $warningCount
    $result.healthy_policies = $healthyCount
    $result.recommendation_policies = $recommendCount
    
    if ($warningCount -gt 0) {{
        $result.status = "HAS_WARNINGS"
        $result.remark = "$warningCount policy warning(s), $recommendCount recommendation(s)"
    }} elseif ($recommendCount -gt 0) {{
        $result.status = "HAS_RECOMMENDATIONS"
        $result.remark = "$recommendCount policy recommendation(s)"
    }} elseif ($policyList.Count -gt 0) {{
        $result.status = "HEALTHY"
        $result.remark = "All $($policyList.Count) policies healthy"
    }} else {{
        $result.status = "NO_POLICIES"
        $result.remark = "No user IRM policies found (only tenant settings)"
    }}

}} catch {{
    $errMsg = $_.Exception.Message
    if ($errMsg -match "access denied|Forbidden|unauthorized|permission") {{
        $result.status = "NO_PERMISSION"
        $result.remark = $errMsg.Substring(0, [Math]::Min(200, $errMsg.Length))
    }} elseif ($errMsg -match "authentication|password|credential|sign-in") {{
        $result.status = "AUTH_FAILED"
        $result.remark = $errMsg.Substring(0, [Math]::Min(200, $errMsg.Length))
    }} else {{
        $result.status = "ERROR"
        $result.remark = $errMsg.Substring(0, [Math]::Min(200, $errMsg.Length))
    }}
}} finally {{
    try {{ Disconnect-ExchangeOnline -Confirm:$false -ErrorAction SilentlyContinue 2>$null }} catch {{}}
}}

$result | ConvertTo-Json -Depth 5 -Compress
'''


def load_accounts():
    accounts = []
    wb = openpyxl.load_workbook(SPN_FILE, read_only=True)
    ws = wb[SPN_SHEET]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    user_col = headers.index("odluser")
    pwd_col = headers.index("odlpassword")
    tid_col = headers.index("TenantId")
    for row in ws.iter_rows(min_row=2, values_only=True):
        username = row[user_col]
        password = row[pwd_col]
        tenant_id = row[tid_col]
        if username and password and tenant_id:
            accounts.append({
                "username": str(username).strip(),
                "password": str(password).strip(),
                "tenant_id": str(tenant_id).strip(),
            })
    wb.close()
    return accounts


def check_account(account, index, total):
    """Check IRM policies for a single account via PowerShell."""
    username = account["username"]
    tenant_id = account["tenant_id"]
    password = account["password"]
    
    print(f"\n{'='*60}")
    print(f"[{index}/{total}] {username}")
    print(f"  TenantID: {tenant_id}")
    
    result = {
        "username": username,
        "tenant_id": tenant_id,
        "irm_status": "ERROR",
        "total_policies": "",
        "warning_policies": "",
        "recommendation_policies": "",
        "healthy_policies": "",
        "policy_details": "",
        "remark": "",
        "checked_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }
    
    # Build PS script with escaped credentials
    ps_script = PS_SCRIPT_TEMPLATE.format(
        username=username.replace('"', '`"'),
        password=password.replace('"', '`"'),
    )
    
    try:
        proc = subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps_script],
            capture_output=True, text=True, timeout=180,
            encoding="utf-8", errors="replace"
        )
        
        output = proc.stdout.strip()
        stderr = proc.stderr.strip()
        
        if stderr:
            # Filter out known noise
            noise = ["WARNING", "we have made updates", "REST-based", "deprecating"]
            relevant_err = [l for l in stderr.split("\n") if not any(n.lower() in l.lower() for n in noise)]
            if relevant_err:
                print(f"  STDERR: {' '.join(relevant_err)[:200]}")
        
        # Find JSON in output (last JSON block)
        json_data = None
        for line in reversed(output.split("\n")):
            line = line.strip()
            if line.startswith("{"):
                try:
                    json_data = json.loads(line)
                    break
                except json.JSONDecodeError:
                    continue
        
        if json_data:
            result["irm_status"] = json_data.get("status", "ERROR")
            result["total_policies"] = str(json_data.get("total_policies", ""))
            result["warning_policies"] = str(json_data.get("warning_policies", ""))
            result["recommendation_policies"] = str(json_data.get("recommendation_policies", ""))
            result["healthy_policies"] = str(json_data.get("healthy_policies", ""))
            result["remark"] = json_data.get("remark", "")
            
            # Format policy details
            policies = json_data.get("policies", [])
            if isinstance(policies, dict):
                policies = [policies]
            if policies:
                details = []
                for p in policies:
                    name = p.get("name", "?")
                    health = p.get("health_status", "?")
                    ind = f"{p.get('indicators_enabled', 0)}/{p.get('indicators_total', 0)}"
                    detail = f"{name} [health:{health}, ind:{ind}]"
                    if p.get("has_warning"):
                        vd = p.get("validation_details", "")
                        detail += f" WARNING"
                        if vd:
                            detail += f":{vd}"
                    details.append(detail)
                result["policy_details"] = "; ".join(details)
                
                print(f"  Policies: {len(policies)}")
                for p in policies:
                    health = p.get("health_status", "?")
                    warn_flag = " *** WARNING ***" if p.get("has_warning") else ""
                    rec_flag = f" (recs:{p.get('recommend_count', 0)})" if p.get("recommend_count", 0) > 0 else ""
                    print(f"    - {p.get('name')}: {health} | ind:{p.get('indicators_enabled')}/{p.get('indicators_total')}{warn_flag}{rec_flag}")
        else:
            result["remark"] = f"No JSON output. stdout: {output[:200]}"
            print(f"  Could not parse output: {output[:200]}")
            
    except subprocess.TimeoutExpired:
        result["irm_status"] = "TIMEOUT"
        result["remark"] = "PowerShell timed out after 180s"
        print(f"  TIMEOUT after 180s")
    except Exception as e:
        result["remark"] = str(e)[:200]
        print(f"  ERROR: {e}")
    
    print(f"  RESULT: {result['irm_status']} | {result['remark'][:60]}")
    return result


def save_results(results):
    fieldnames = ["username", "tenant_id", "irm_status", "total_policies",
                  "warning_policies", "recommendation_policies", "healthy_policies",
                  "policy_details", "remark", "checked_at"]
    with open(RESULTS_FILE, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(results)


def main():
    accounts = load_accounts()
    print(f"Loaded {len(accounts)} accounts from sheet '{SPN_SHEET}'")
    
    if START_INDEX > 0:
        accounts = accounts[START_INDEX:]
    if MAX_COUNT > 0:
        accounts = accounts[:MAX_COUNT]
    
    total = len(accounts)
    print(f"Checking {total} accounts for IRM policy warnings via PowerShell CLI")
    print(f"Output: {RESULTS_FILE}")
    
    results = []
    for i, acc in enumerate(accounts, 1):
        r = check_account(acc, i, total)
        results.append(r)
        save_results(results)
    
    # Summary
    print(f"\n{'='*60}")
    print(f"IRM CLI CHECK SUMMARY - {SPN_SHEET}")
    print(f"{'='*60}")
    from collections import Counter
    status_counts = Counter(r["irm_status"] for r in results)
    for s, c in sorted(status_counts.items()):
        print(f"  {s}: {c}")
    print(f"  TOTAL: {total}")
    
    warned = [r for r in results if r["irm_status"] == "HAS_WARNINGS"]
    if warned:
        print(f"\n  Accounts with warnings ({len(warned)}):")
        for r in warned:
            print(f"    {r['username']}: {r['warning_policies']} warnings")
            print(f"      {r['policy_details'][:120]}")
    
    print(f"\nResults saved to: {RESULTS_FILE}")


if __name__ == "__main__":
    main()
